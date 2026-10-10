"""The source plugin's lint --selftest, ported to the engine linter on the Diametral pack: a
deliberately off-brand slide where every check fires, plus shapes that must stay silent."""

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from calque_engine.build import build
from calque_engine.lint import DEEP_NESTING, lint
from calque_engine.pack import load_pack

PACK = Path(__file__).resolve().parents[1]
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _box(slide, left, top, width, height, text, font="Geist", color=None, size=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = font
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return tb


@pytest.fixture(scope="module")
def selftest(tmp_path_factory):
    pack = load_pack(PACK)
    tmp = tmp_path_factory.mktemp("selftest")
    spec = {
        "pack_id": "diametral",
        "language": "en",
        "title": "Selftest",
        "slides": [
            {
                "id": "s1",
                "message": "One content slide",
                "message_type": "narrative",
                "form": "content",
                "source": {"kind": "clone", "role": "content", "values": {}},
            }
        ],
    }
    report = build(spec, pack, tmp / "clean.pptx")
    prs = Presentation(str(tmp / "clean.pptx"))
    slide = prs.slides[0]
    red = pack.color("highlight")

    bad = _box(slide, 1, 2, 3, 0.4, "Lorem ipsum XX", font="Comic Sans MS", color="123456")
    for left in (4, 6):  # two highlight elements
        _box(slide, left, 3, 1, 0.4, "actif", color=red)
    _box(slide, 9.6, 1, 2, 0.4, "hors cadre")  # off canvas
    _box(slide, 9.5, 5.3, 0.4, 0.2, "7")  # literal page number, slide is 1
    slop = _box(slide, 1, 1, 4, 0.8, "Robust pipeline — seamless")  # em-dash ERROR + words WARN
    deep = slop.text_frame.add_paragraph()
    deep.level = DEEP_NESTING
    deep.add_run().text = "trop profond"
    over = _box(slide, 1, 4, 1, 0.3, "Une phrase nettement trop longue pour cette boite etroite.")
    over.text_frame.word_wrap = True
    over.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    wide = _box(slide, 5, 4, 1, 0.3, "Un libellé qui déborde sur le côté", size=14)
    wide.text_frame.word_wrap = False
    late = _box(slide, 1, 3.5, 3, 0.4, "visible")
    lp = late.text_frame.paragraphs[0]._p
    lp.append(lp.makeelement(qn("a:endParaRPr"), {}))
    ghost = lp.makeelement(qn("a:r"), {})
    ghost.append(ghost.makeelement(qn("a:t"), {}))
    ghost[0].text = "jamais rendu"
    lp.append(ghost)  # run after endParaRPr
    shout = _box(slide, 1, 0.5, 4, 0.4, "Livré en mars !")  # exclamation mark

    # Must stay SILENT: a price range, an unused off-brand list-style default, the running footer.
    rng = _box(
        slide, 5, 1.5, 3, 0.4, "€15k – €120k / mission, 250€k – €550k", color=pack.color("ink")
    )
    rng.text_frame._txBody.insert(
        1,
        etree.fromstring(
            f'<a:lstStyle xmlns:a="{A}"><a:lvl1pPr><a:defRPr>'
            '<a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
            '<a:latin typeface="Arial"/></a:defRPr></a:lvl1pPr></a:lstStyle>'
        ),
    )
    run_foot = _box(slide, 8.5, 5.43, 1.4, 0.1, "2026 – Diametral – Business presentation")

    # The plugin's table cell carried a middle dot, which this pack deliberately allows.
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(4.5), Inches(3), Inches(0.4))
    table.table.cell(0, 0).text = "Lot 1 — Lot 2"
    slide.notes_slide.notes_text_frame.text = "Contexte – enjeux"

    path = tmp / "selftest.pptx"
    prs.save(str(path))
    findings = lint(path, pack, "en", dict(report.slides.values()))
    ids = {
        "bad": bad.shape_id,
        "slop": slop.shape_id,
        "table": table.shape_id,
        "shout": shout.shape_id,
        "silent": {rng.shape_id, run_foot.shape_id},
        "path": path,
        "tmap": dict(report.slides.values()),
    }
    return findings, ids


def test_every_check_fires(selftest):
    findings, _ = selftest
    fired = {f.check for f in findings}
    expected = {"font", "palette", "single-use", "placeholder", "off-canvas", "page-number"}
    expected |= {"overflow", "slop", "nesting", "run-order"}
    assert expected <= fired, findings
    assert any("non-wrapping" in f.message for f in findings)


def test_slop_tiers_and_reach(selftest):
    findings, ids = selftest
    slop = [f for f in findings if f.check == "slop"]
    tiers = {(f.shape_id, f.severity) for f in slop}
    assert (ids["slop"], "ERROR") in tiers  # em-dash: pack rule
    assert (ids["slop"], "WARN") in tiers  # robust / seamless: core vocabulary, en
    assert (ids["shout"], "ERROR") in tiers  # exclamation mark: pack rule
    assert any(f.shape_id == ids["table"] and "'—'" in f.message for f in slop)
    assert any(
        f.severity == "NOTE" and "speaker notes" in f.message and "'–'" in f.message for f in slop
    )


def test_single_highlight_rule(selftest):
    findings, _ = selftest
    (red,) = [f for f in findings if f.check == "single-use"]
    assert red.severity == "WARN" and "on 2 shapes" in red.message


def test_no_false_positive(selftest):
    findings, ids = selftest
    assert [str(f) for f in findings if f.shape_id in ids["silent"]] == []


def test_vocabulary_is_language_scoped(selftest):
    findings, ids = selftest
    fr = lint(ids["path"], load_pack(PACK), "fr", ids["tmap"])
    tiers = {(f.shape_id, f.severity) for f in fr if f.check == "slop"}
    assert (ids["slop"], "ERROR") in tiers  # em-dash: lang any
    assert (ids["slop"], "WARN") not in tiers  # robust / seamless: en only


def test_closing_slide_keeps_its_signature_line(tmp_path):
    pack = load_pack(PACK)
    spec = {
        "pack_id": "diametral",
        "language": "en",
        "title": "Closing",
        "slides": [
            {
                "id": "end",
                "message": "Close",
                "message_type": "closing",
                "form": "closing",
                "source": {"kind": "clone", "role": "closing", "values": {}},
            }
        ],
    }
    report = build(spec, pack, tmp_path / "c.pptx")
    tmap = dict(report.slides.values())
    assert [f for f in lint(tmp_path / "c.pptx", pack, "en", tmap) if f.check == "slop"] == []
    # without a map, an unchanged closing clone is still recognised by its text
    assert [f for f in lint(tmp_path / "c.pptx", pack, "en") if f.check == "slop"] == []
    # rewritten, its copy is the author's again: no exemption without a map
    spec["slides"][0]["source"]["values"] = {"1710": "See you soon!"}
    build(spec, pack, tmp_path / "d.pptx")
    unexempt = lint(tmp_path / "d.pptx", pack, "en")
    assert any(f.check == "slop" and "'!'" in f.message for f in unexempt)
    assert not [
        f
        for f in lint(tmp_path / "c.pptx", pack, "en", exempt_closing_slides={1})
        if f.check == "slop"
    ]


def test_closing_line_in_the_deck_language(tmp_path):
    """A French deck closes in French (pack.yaml `localized_text`), still exempt as a signature."""
    pack = load_pack(PACK)
    end = {"kind": "clone", "role": "closing", "values": {}}
    slide = {"id": "end", "message": "Close", "message_type": "closing", "form": "closing"}
    for language, line in (("fr", "Merci de votre attention."), ("en", "Thanks for watching!")):
        spec = {"pack_id": "diametral", "language": language, "title": "Fin", "slides": []}
        spec["slides"].append({**slide, "source": end})
        out = tmp_path / f"{language}.pptx"
        tmap = dict(build(spec, pack, out).slides.values())
        (shape,) = [s for s in Presentation(str(out)).slides[0].shapes if s.shape_id == 1710]
        assert " ".join(shape.text_frame.text.split()) == line
        for m in (tmap, None):
            assert [f for f in lint(out, pack, language, m) if f.check == "slop"] == []


def test_webinar_cover_and_closing(tmp_path):
    """The persona-review webinar deck: a keyword too long for its tag, a closing rewritten in
    off-charter copy. Built (server, with the map) and linted bare (CLI), the findings agree."""
    pack = load_pack(PACK)

    def deck(title):
        cover = {"26": title, "27": "TRANSFORMATION ACHATS", "30": "WEBINAR", "31": "10/2026"}
        closing = {"1710": "Réservez un rendez-vous — c'est révolutionnaire !"}
        spec = {"pack_id": "diametral", "language": "fr", "title": "Webinar", "slides": []}
        for sid, role, values in (("c", "cover", cover), ("e", "closing", closing)):
            src = {"kind": "clone", "role": role, "values": values}
            spec["slides"].append(
                {"id": sid, "message": sid, "message_type": role, "form": role, "source": src}
            )
        report = build(spec, pack, tmp_path / "w.pptx")
        mapped = lint(tmp_path / "w.pptx", pack, "fr", dict(report.slides.values()))
        assert mapped == lint(tmp_path / "w.pptx", pack, "fr")
        return {(f.slide, f.shape_id, f.check, f.severity) for f in mapped}

    found = deck("L'IA dans les achats\u202f: ce que mesurent les CPO")
    assert (1, 27, "capacity", "WARN") in found  # the grown keyword tag runs into the date
    assert {(2, 1710, "slop", "ERROR"), (2, 1710, "slop", "WARN")} <= found
    # The title holds two lines of Ufficio at 52 pt: its three-line preview came from a fallback
    # face, which lint reports when the pack fonts are missing.
    assert not {f for f in found if f[1] == 26}
    if not (PACK / "fonts" / "Ufficio-300.otf").is_file():
        assert (None, None, "fonts", "WARN") in found
    long = deck(
        "L'IA dans les achats\u202f: ce que mesurent vraiment les directions achats en 2026"
    )
    assert (1, 26, "capacity", "WARN") in long


def test_voice_hard_rules(tmp_path):
    """voice.md naming and mechanics, enforced from pack.yaml: brand spelling, AI-native, etc.,
    filler buzzwords; every banned word of a shape is reported, not only the first."""
    pack = load_pack(PACK)
    spec = {
        "pack_id": "diametral",
        "language": "en",
        "title": "Voice",
        "slides": [
            {
                "id": "s1",
                "message": "Voice",
                "message_type": "narrative",
                "form": "content",
                "source": {"kind": "clone", "role": "content", "values": {}},
            }
        ],
    }
    report = build(spec, pack, tmp_path / "v.pptx")
    prs = Presentation(str(tmp_path / "v.pptx"))
    slide = prs.slides[0]
    copy = (
        "Diam\u00e9tral and DIAMETRAL are AI-First, AI-Native, an Agile 360 solution, data, "
        "models, etc. AI-powered, scalable and disruptif."
    )
    bad = _box(slide, 0.4, 1.6, 7, 1.2, copy)
    good = _box(slide, 0.4, 3, 7, 0.6, "Diametral is AI-native: write to contact@diametral.com.")
    prs.save(str(tmp_path / "voice.pptx"))
    found = lint(tmp_path / "voice.pptx", pack, "en", dict(report.slides.values()))
    slop = [(f.severity, f.message.split(":")[0]) for f in found if f.shape_id == bad.shape_id]
    for sev, hit in [
        ("ERROR", "'Diam\u00e9tral'"),
        ("ERROR", "'DIAMETRAL'"),
        ("ERROR", "'AI-First'"),
        ("ERROR", "'AI-Native'"),
        ("ERROR", "'etc.'"),
        ("WARN", "'Agile'"),
        ("WARN", "'360 solution'"),
        ("WARN", "'AI-powered'"),
        ("WARN", "'scalable'"),
        ("WARN", "'disruptif'"),
    ]:
        assert (sev, hit) in slop, slop
    assert not [str(f) for f in found if f.shape_id == good.shape_id and f.check == "slop"]
