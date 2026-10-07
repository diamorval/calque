"""Phase 2 verify: a title edited on an imported deck keeps its formatting."""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from calque_engine.build import build
from calque_engine.importer import import_pptx
from calque_engine.pack import load_pack
from calque_engine.patch import patch

PACK = Path(__file__).resolve().parents[1]
DECK = PACK / "reference" / "business-presentation.pptx"


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def test_title_edit_on_imported_deck_keeps_formatting(tmp_path):
    pack = load_pack(PACK)
    deck, tmap = import_pptx(DECK, pack, tmp_path, language="fr")
    original = Presentation(str(DECK))
    n = 2
    slide = original.slides[n - 1]
    shape = max(
        (sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()),
        key=lambda sh: max(
            [r.font.size or 0 for p in sh.text_frame.paragraphs for r in p.runs] or [0]
        ),
    )
    rpr_before = etree.tostring(shape._element.find(".//" + qn("a:rPr")))
    ppr = shape.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    ppr_before = etree.tostring(ppr) if ppr is not None else b""

    patched, _ = patch(
        deck,
        [{"op": "set", "slide": f"s{n}", "shape_id": shape.shape_id, "value": "Un titre modifié"}],
        pack,
    )
    out = tmp_path / "out.pptx"
    build(patched, pack, out, base=tmp_path / deck["base"])

    result = Presentation(str(out))
    assert len(result.slides) == len(original.slides)
    edited = next(sh for sh in result.slides[n - 1].shapes if sh.shape_id == shape.shape_id)
    assert edited.text_frame.text == "Un titre modifié"
    assert etree.tostring(edited._element.find(".//" + qn("a:rPr"))) == rpr_before
    ppr = edited.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    assert (etree.tostring(ppr) if ppr is not None else b"") == ppr_before
    for i, (a, b) in enumerate(zip(original.slides, result.slides, strict=True), start=1):
        if i != n:
            assert _texts(a) == _texts(b), f"slide {i} changed"
