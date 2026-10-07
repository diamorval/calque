"""Builds template.pptx for the acme-test pack: python-pptx's stock theme, 13.33 x 7.5 in, four
role slides. A deliberately plain second pack: proves the engine needs nothing brand-specific.

    uv run --project ../../engine python make_template.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).parent


def page_number(slide, text: str) -> None:
    tb = slide.shapes.add_textbox(Inches(12.3), Inches(6.95), Inches(0.6), Inches(0.3))
    tb.text_frame.text = text
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(10)


def place(shape, l: float, t: float, w: float, h: float) -> None:
    shape.left, shape.top, shape.width, shape.height = Inches(l), Inches(t), Inches(w), Inches(h)


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title_slide, section, title_only = (prs.slide_layouts[i] for i in (0, 2, 5))

    s = prs.slides.add_slide(title_slide)
    s.shapes.title.text = "Presentation title"
    s.placeholders[1].text = "Subtitle and date"
    place(s.shapes.title, 0.8, 2.3, 11.7, 1.6)
    place(s.placeholders[1], 0.8, 4.1, 11.7, 1.0)

    s = prs.slides.add_slide(section)
    s.shapes.title.text = "Section title"
    s.placeholders[1].text = "Section subtitle"
    place(s.shapes.title, 0.8, 2.9, 11.7, 1.4)
    place(s.placeholders[1], 0.8, 4.4, 11.7, 0.8)
    page_number(s, "2")

    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Slide title"
    place(s.shapes.title, 0.6, 0.35, 12.1, 1.0)
    page_number(s, "3")

    s = prs.slides.add_slide(title_slide)
    s.shapes.title.text = "Contact"
    s.placeholders[1].text = "Name SURNAME, email"
    place(s.shapes.title, 0.8, 2.3, 11.7, 1.6)
    place(s.placeholders[1], 0.8, 4.1, 11.7, 1.0)

    prs.save(HERE / "template.pptx")


if __name__ == "__main__":
    main()
