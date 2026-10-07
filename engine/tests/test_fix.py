"""Safe fixes: ten defects planted in a clean deck are all found by lint, and gone after fix."""

import copy
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor

from calque_engine.fix import fix, nearest
from calque_engine.lint import lint

from .test_lint import POS, _box, clean, errors, f_font, f_off_canvas, f_palette  # noqa: F401


def plant(src: Path, pack, out: Path) -> int:
    """Ten safe-fixable defects: 2 palette, 3 font, 2 off-canvas, nesting, run-order, theme."""
    prs = Presentation(str(src))
    w, h = prs.slide_width / 914400, prs.slide_height / 914400
    s = prs.slides[POS - 1]
    f_palette(s, pack, w, h)
    greys = [g for g in (f"{i:02X}" * 3 for i in range(1, 255)) if g not in pack.palette()]
    _box(s, pack, 1, 3, 3, 0.4, "fill").font.color.rgb = RGBColor.from_string(greys[40])
    f_font(s, pack, w, h)
    _box(s, pack, 1, 3.5, 3, 0.4, "face 2").font.name = "Comic Sans MS"
    _box(prs.slides[0], pack, 1, 4, 3, 0.4, "face 3").font.name = "Papyrus"
    f_off_canvas(s, pack, w, h)
    _box(s, pack, 1, -1, 2, 0.4, "above")
    _box(s, pack, 5, 2, 3, 1, "top")
    shape = s.shapes[-1]
    p2 = shape.text_frame.add_paragraph()
    p2.text, p2.level = "deep", 2
    p = shape.text_frame.paragraphs[0]._p
    end = p.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr", {})
    p.append(end)
    late = copy.deepcopy(p.find("{http://schemas.openxmlformats.org/drawingml/2006/main}r"))
    p.append(late)
    prs.save(str(out))
    with zipfile.ZipFile(out) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if re.match(r"ppt/theme/theme1\.xml$", info.filename):
                data = re.sub(
                    rb'(<a:majorFont><a:latin typeface=")[^"]*', rb"\1Wingdings", data, count=1
                )
            z.writestr(info, data)
    return 10


def test_nearest():
    red, blue = f"{255:02X}0000", f"0000{255:02X}"
    assert nearest(f"{254:02X}0101", {red, blue}) == red


def test_ten_defects_found_and_fixed(clean, tmp_path):  # noqa: F811
    pack, path, tmap = clean
    bad = tmp_path / "bad.pptx"
    plant(path, pack, bad)
    found = [f for f in lint(bad, pack, "en", tmap) if f.severity == "ERROR"]
    checks = sorted(f.check for f in found)
    assert checks == sorted(
        ["palette"] * 2 + ["font"] * 3 + ["off-canvas"] * 2 + ["nesting", "run-order", "theme"]
    ), found
    applied = fix(bad, pack, tmp_path / "fixed.pptx")
    assert {a["check"] for a in applied} == set(checks)
    assert errors(lint(tmp_path / "fixed.pptx", pack, "en", tmap)) == []
