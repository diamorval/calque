"""A built deck's document properties are its own: no template authors, edit history or library
defaults."""

import zipfile
from datetime import UTC, datetime

from pptx import Presentation
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

from calque_engine.build import build
from calque_engine.pack import load_pack

CHANGES = "http://schemas.microsoft.com/office/2016/11/relationships/changesInfo"
CHANGES_CT = "application/vnd.ms-powerpoint.changesinfo+xml"


def _deck():
    return {
        "pack_id": "neutral",
        "language": "en",
        "title": "Quarterly review",
        "slides": [
            {
                "id": "a",
                "message": "Cover",
                "message_type": "cover",
                "form": "cover",
                "source": {"kind": "clone", "role": "cover", "values": {"2": "A title"}},
            }
        ],
    }


def _leaky_template(pack_dir):
    """The template as an office suite leaves it: its authors, and their revision history."""
    path = pack_dir / "template.pptx"
    prs = Presentation(str(path))
    prs.core_properties.author = "Template Writer"
    prs.core_properties.last_modified_by = "Template Editor"
    prs.core_properties.revision = 42
    history = b'<chgInfo xmlns="urn:x">Template Reviewer changed slide 2</chgInfo>'
    part = Part(
        PackURI("/ppt/changesInfos/changesInfo1.xml"), CHANGES_CT, prs.part.package, history
    )
    prs.part.relate_to(part, CHANGES)
    prs.save(str(path))


def _strings(pptx) -> str:
    with zipfile.ZipFile(pptx) as z:
        return "\n".join(f"{n}\n{z.read(n).decode('utf-8', 'replace')}" for n in z.namelist())


def test_built_deck_names_no_template_author(neutral_pack, tmp_path):
    _leaky_template(neutral_pack)
    assert "Template Reviewer" in _strings(neutral_pack / "template.pptx")
    out = tmp_path / "deck.pptx"
    build(_deck(), load_pack(neutral_pack), out)

    text = _strings(out)
    for leak in (
        "Template Writer",
        "Template Editor",
        "Template Reviewer",
        "changesInfo",
        "python-pptx",
        "Steve Canny",
    ):
        assert leak not in text
    core = Presentation(str(out)).core_properties  # still opens
    assert core.title == "Quarterly review"
    assert (core.author, core.last_modified_by, core.revision) == ("", "", 1)
    now = datetime.now(UTC).replace(tzinfo=None)
    assert core.created == core.modified and abs((now - core.created).total_seconds()) < 600


def test_author_when_passed(neutral_pack, tmp_path):
    out = tmp_path / "deck.pptx"
    build(_deck(), load_pack(neutral_pack), out, author="Alice Martin")
    core = Presentation(str(out)).core_properties
    assert (core.author, core.last_modified_by) == ("Alice Martin", "Alice Martin")
