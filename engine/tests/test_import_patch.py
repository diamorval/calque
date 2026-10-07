import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from calque_engine.build import build
from calque_engine.importer import import_pptx
from calque_engine.pack import load_pack
from calque_engine.patch import PatchError, apply_ops, patch


def _deck():
    return {
        "pack_id": "neutral",
        "language": "en",
        "title": "T",
        "slides": [
            {
                "id": "a",
                "message": "Cover",
                "message_type": "cover",
                "form": "cover",
                "source": {"kind": "clone", "role": "cover", "values": {"2": "Original title"}},
            },
            {
                "id": "b",
                "message": "Body",
                "message_type": "narrative",
                "form": "content",
                "source": {"kind": "clone", "role": "content", "values": {"2": "A slide"}},
                "notes": "Keep these notes.",
            },
        ],
    }


def _rpr(shape) -> bytes:
    r = shape._element.find(".//" + qn("a:r"))
    rpr = r.find(qn("a:rPr"))
    return etree.tostring(rpr) if rpr is not None else b""


def test_import_then_patch_keeps_formatting_and_notes(neutral_pack, tmp_path):
    pack = load_pack(neutral_pack)
    src = tmp_path / "src.pptx"
    build(_deck(), pack, src)
    # give the title a distinctive run format the patch must keep
    prs = Presentation(str(src))
    title = prs.slides[0].shapes.title
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    title.text_frame.paragraphs[0].runs[0].font.italic = True
    prs.save(str(src))
    before = _rpr(Presentation(str(src)).slides[0].shapes.title)

    deck, tmap = import_pptx(src, pack, tmp_path / "work", language="en")
    assert [s["form"] for s in deck["slides"]] == ["imported", "imported"]
    assert deck["slides"][0]["message"] == "Original title"
    title_id = tmap["slides"][0]["shapes"][0]["id"]

    patched, _ = patch(
        deck, [{"op": "set", "slide": "s1", "shape_id": title_id, "value": "New title"}], pack
    )
    out = tmp_path / "out.pptx"
    build(patched, pack, out, base=tmp_path / "work" / deck["base"])

    after = Presentation(str(out))
    assert after.slides[0].shapes.title.text_frame.text == "New title"
    assert _rpr(after.slides[0].shapes.title) == before
    assert after.slides[1].notes_slide.notes_text_frame.text == "Keep these notes."
    assert len(after.slides) == 2


def test_patch_ops():
    d = _deck()
    out = apply_ops(
        d,
        [
            {"op": "set_field", "slide": "b", "field": "title", "value": "X"},
            {"op": "move_slide", "slide": "b", "to": 0},
            {"op": "delete_slide", "slide": "a"},
        ],
    )
    assert [s["id"] for s in out["slides"]] == ["b"] and out["slides"][0]["title"] == "X"
    assert d["slides"][0]["id"] == "a"  # input untouched
    with pytest.raises(PatchError, match="no slide 'zz'"):
        apply_ops(d, [{"op": "delete_slide", "slide": "zz"}])
    with pytest.raises(PatchError, match="set_params applies"):
        apply_ops(d, [{"op": "set_params", "slide": "a", "params": {}}])


def test_style_only_set_keeps_text():
    def set_(d, value):
        return apply_ops(d, [{"op": "set", "slide": "a", "shape_id": 2, "value": value}])

    out = set_(_deck(), {"color": "accent"})
    assert out["slides"][0]["source"]["values"]["2"] == {
        "text": "Original title",
        "color": "accent",
    }
    out = set_(out, {"bold": True})
    assert out["slides"][0]["source"]["values"]["2"] == {
        "text": "Original title",
        "color": "accent",
        "bold": True,
    }
    assert set_(_deck(), "New")["slides"][0]["source"]["values"]["2"] == "New"
