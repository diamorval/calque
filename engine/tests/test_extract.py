import json
import zipfile

import pytest
import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor

from calque_engine import tokens as tk
from calque_engine.api import op_extract, op_lint
from calque_engine.extract import _role_numbers, capacity, draft_tokens, extract

from .conftest import make_template


def test_extract_neutral_template(tmp_path):
    tmap = extract(make_template(tmp_path / "t.pptx"))
    assert [s["number"] for s in tmap["slides"]] == [1, 2, 3]
    assert set(tmap["theme"]["colors"]) >= {"dk1", "lt1", "accent1", "hlink"}
    assert tmap["theme"]["fonts"]["major"]
    title = next(sh for sh in tmap["slides"][1]["shapes"] if sh.get("text") == "Content title")
    assert isinstance(title["id"], int) and len(title["bbox"]) == 4


def test_draft_tokens_cover_every_required_role(tmp_path):
    values = tk.resolve(draft_tokens(extract(make_template(tmp_path / "t.pptx"))))
    assert tk.missing_roles(values) == []


def test_reviewed_keys_survive_reextraction(tmp_path):
    path = make_template(tmp_path / "t.pptx")
    first = extract(path)
    first["slides"][0]["slots"] = {"title": 2}
    first["slides"][0]["description"] = "Cover"
    again = extract(path, previous=first)
    assert again["slides"][0]["slots"] == {"title": 2}
    assert again["slides"][0]["description"] == "Cover"


def test_capacity():
    assert capacity(4.0, 1.0, 12) == {"chars_per_line": 48, "lines": 5}


def test_draft_manifest_makes_a_valid_pack(tmp_path):
    from calque_engine.extract import draft_manifest
    from calque_engine.pack import load_pack

    d = tmp_path / "p"
    d.mkdir()
    tmap = extract(make_template(d / "template.pptx"))
    manifest = draft_manifest(tmap, "p", "P")
    assert manifest["roles"] == {"cover": [1], "content": [2], "closing": [3]}
    assert tmap["slides"][0]["slots"]["title"]
    assert "Deck title" in manifest["lint"]["placeholders"]
    (d / "pack.yaml").write_text(yaml.safe_dump(manifest, allow_unicode=True))
    (d / "template-map.yaml").write_text(yaml.safe_dump(tmap, sort_keys=False))
    (d / "tokens.json").write_text(json.dumps(draft_tokens(tmap)))
    assert load_pack(d).id == "p"


def test_potx_is_rewritten_as_a_presentation(tmp_path):
    potx = tmp_path / "t.potx"
    with (
        zipfile.ZipFile(make_template(tmp_path / "t.pptx")) as z,
        zipfile.ZipFile(potx, "w") as out,
    ):
        for item in z.infolist():
            data = z.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentation.main+xml", b"template.main+xml")
            out.writestr(item, data)
    with pytest.raises(ValueError):
        Presentation(str(potx))
    assert op_extract(str(potx))["manifest"]["roles"]["cover"] == [1]
    assert len(Presentation(str(potx)).slides) == 3


def test_fonts_and_colours_set_on_runs_are_declared(tmp_path):
    d = tmp_path / "p"
    d.mkdir()
    path = make_template(d / "template.pptx")
    prs = Presentation(str(path))
    run = prs.slides[1].shapes.title.text_frame.paragraphs[0].runs[0]
    run.font.name = "Run Face"
    color = RGBColor(18, 52, 86)  # not in the default theme
    run.font.color.rgb = color
    prs.save(str(path))

    r = op_extract(str(path), "p", "P")
    assert r["manifest"]["lint"]["extra_fonts"] == ["Run Face"]
    assert r["manifest"]["lint"]["extra_colors"] == [str(color)]
    assert r["review"]["fonts"]["role.font.body"] == r["template_map"]["theme"]["fonts"]["minor"]
    (d / "pack.yaml").write_text(yaml.safe_dump(r["manifest"], allow_unicode=True))
    (d / "template-map.yaml").write_text(yaml.safe_dump(r["template_map"]))
    (d / "tokens.json").write_text(json.dumps(r["tokens"]))
    findings = op_lint(str(d), str(path), template=True)["findings"]
    assert [f for f in findings if f["severity"] == "ERROR"] == []


def test_roles_from_text_structure():
    def slide(n, *texts):
        shapes = [
            {"id": i, "kind": "text", "text": t, "size": s, "bbox": [0.5, y, 9, 1]}
            for i, (t, s, y) in enumerate(texts, start=2)
        ]
        return {"number": n, "layout": "Custom", "page_number": None, "shapes": shapes}

    slides = [
        slide(1, ("Deck title", 40, 2)),
        slide(2, ("Agenda", 32, 0.3), ("Context", 16, 2), ("Plan", 16, 3)),
        slide(3, ("Context", 40, 2), ("01", 32, 1)),
        slide(4, ("A long action title for the content", 24, 0.3), ("Body", 12, 2)),
        slide(5, ("Plan", 40, 2), ("02", 32, 1)),
        slide(6, ("Thank you", 40, 2)),
        slide(7, ("Appendix notes", 24, 0.3), ("More", 12, 2), ("More", 12, 3), ("x", 9, 4)),
    ]
    assert _role_numbers(slides, 7.5) == {
        "cover": [1],
        "summary": [2],
        "divider": [3, 5],
        "content": [4],
        "closing": [6],
    }
