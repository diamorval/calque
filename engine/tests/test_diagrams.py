"""Every diagram on every pack: builds through the DeckSpec, stays on the canvas and above the
footer, highlights at most one element."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from calque_engine.build import build
from calque_engine.pack import load_pack
from calque_engine.style import Style

REPO = Path(__file__).resolve().parents[2]
PACKS = sorted((REPO / "packs").glob("*/pack.yaml"))

STEPS = [
    {"label": "Frame", "sublabel": "3 weeks"},
    {"label": "Migrate", "sublabel": "8 weeks"},
    {"label": "Test", "sublabel": "4 weeks"},
    {"label": "Switch", "sublabel": "1 week"},
    {"label": "Run"},
]

# id -> (message_type, params, expected minimum of drawn shapes)
DIAGRAMS = {
    "flow": ("process", {"steps": STEPS, "active": 1}, 14),
    "swimlane": (
        "process",
        {
            "lanes": [
                {"label": "Client", "steps": ["Brief", "Review", "Sign"]},
                {"label": "Team", "steps": ["Audit", {"label": "Build", "sublabel": "6 wk"}]},
                {"label": "Ops", "steps": ["Deploy", "Run", "Report"]},
            ],
            "active": [1, 1],
        },
        11,
    ),
    "layers": (
        "system",
        {
            "layers": [
                {"label": "Exposure", "sublabel": "APIs, dashboards"},
                {"label": "Services", "sublabel": "catalogue, quality"},
                {"label": "Processing", "sublabel": "batch, streaming"},
                {"label": "Foundation", "sublabel": "lakehouse, IAM"},
            ],
            "active": 1,
        },
        4,
    ),
    "hub": (
        "system",
        {
            "center": {"label": "Data platform", "sublabel": "one source of truth"},
            "spokes": ["Finance", "Sales", "Supply", "HR", "Marketing", "Risk"],
            "active": 2,
        },
        13,
    ),
    "matrix2x2": (
        "positioning",
        {
            "x_axis": ["Niche", "Global"],
            "y_axis": ["Follower", "Leader"],
            "quadrants": ["Specialists", "Leaders", "Locals", "Generalists"],
            "items": [
                {"label": "Northwind", "x": 0.8, "y": 0.75},
                {"label": "Us", "x": 0.62, "y": 0.86},
                {"label": "Contoso", "x": 0.25, "y": 0.3},
                {"label": "Fabrikam", "x": 0.9, "y": 0.2},
            ],
            "active": 1,
        },
        18,
    ),
    "funnel": (
        "conversion",
        {
            "stages": ["Leads", {"label": "Qualified", "sublabel": "35 %"}, "Proposal", "Signed"],
            "values": [1240, 434, 208, 77],
            "active": 3,
        },
        8,
    ),
    "cycle": (
        "loop",
        {
            "steps": ["Plan", "Build", {"label": "Measure", "sublabel": "weekly"}, "Learn"],
            "active": 2,
        },
        8,
    ),
    "before_after": (
        "transformation",
        {
            "before": {"title": "Today", "lines": ["12 tools, no owner", "Monthly reporting"]},
            "after": {"title": "Target", "lines": ["One platform", "Daily reporting"]},
            "arrow_label": "6 months",
        },
        6,
    ),
}


def deck(pack, form, message_type, source):
    return {
        "pack_id": pack.id,
        "language": "en",
        "title": "Test",
        "slides": [
            {
                "id": "s1",
                "message": "The point of the slide.",
                "message_type": message_type,
                "form": form,
                "title": "The verdict in one line",
                "source": source,
            }
        ],
    }


def drawn(pack, path: Path):
    """Shapes the renderer added: everything not on the pack's content template slide."""
    n = pack.slides_for("content")[0]
    tmpl = {s["number"]: s for s in pack.template_map["slides"]}[n]
    ids = {sh["id"] for sh in tmpl["shapes"]}
    slide = Presentation(str(path)).slides[0]
    return [sh for sh in slide.shapes if sh.shape_id not in ids]


def colors_of(sh) -> set[str]:
    """Colours a shape uses: line, solid fill, text runs (table cells included)."""
    out = set()
    for get in (lambda: sh.line.color.rgb, lambda: sh.fill.fore_color.rgb):
        try:
            out.add(str(get()))
        except (AttributeError, TypeError, NotImplementedError):
            pass
    frames = [sh.text_frame] if getattr(sh, "has_text_frame", False) else []
    if getattr(sh, "has_table", False):
        frames += [c.text_frame for row in sh.table.rows for c in row.cells]
    for tf in frames:
        for p in tf.paragraphs:
            for r in p.runs:
                try:
                    out.add(str(r.font.color.rgb))
                except AttributeError:
                    pass
    return out


def check(pack, path: Path, minimum: int, active: bool) -> list:
    st = Style(pack)
    g = st.grid
    shapes = drawn(pack, path)
    assert len(shapes) >= minimum, f"{len(shapes)} drawn shapes, expected >= {minimum}"
    eps = 0.01
    for sh in shapes:
        l, t = Emu(sh.left).inches, Emu(sh.top).inches
        r, b = l + Emu(sh.width).inches, t + Emu(sh.height).inches
        assert l >= -eps and t >= -eps and r <= g.width + eps, (sh.name, l, t, r)
        assert b <= g.bottom + eps, f"{sh.name} bottom {b:.2f} crosses the footer {g.bottom}"
    hi = pack.color("highlight")
    hot = [sh for sh in shapes if hi in colors_of(sh)]
    if not active:
        assert not hot
    else:
        # one element: its stroke plus its numeral, on the same row or column
        assert 1 <= len(hot) <= 2, [sh.name for sh in hot]
        if len(hot) == 2:
            a, c = hot
            gap_x = max(a.left, c.left) - min(a.left + a.width, c.left + c.width)
            gap_y = max(a.top, c.top) - min(a.top + a.height, c.top + c.height)
            assert min(gap_x, gap_y) <= 0, "two separate highlighted elements"
    return shapes


@pytest.mark.parametrize("manifest", PACKS, ids=lambda p: p.parent.name)
@pytest.mark.parametrize("did", sorted(DIAGRAMS))
def test_diagram_builds_on_grid(manifest, did, tmp_path):
    pack = load_pack(manifest.parent)
    mtype, params, minimum = DIAGRAMS[did]
    out = tmp_path / f"{did}.pptx"
    build(deck(pack, did, mtype, {"kind": "diagram", "id": did, "params": params}), pack, out)
    check(pack, out, minimum, "active" in params)


@pytest.mark.parametrize("manifest", PACKS, ids=lambda p: p.parent.name)
def test_all_diagrams_in_one_deck(manifest, tmp_path):
    pack = load_pack(manifest.parent)
    data = deck(pack, "flow", "process", {"kind": "diagram", "id": "flow", "params": {}})
    data["slides"] = [
        {
            "id": did,
            "message": "m",
            "message_type": mtype,
            "form": did,
            "title": "Verdict",
            "source": {"kind": "diagram", "id": did, "params": params},
        }
        for did, (mtype, params, _) in DIAGRAMS.items()
    ]
    report = build(data, pack, tmp_path / "all.pptx")
    assert len(Presentation(str(report.path)).slides) == len(DIAGRAMS)
