"""Charts are native, role-styled and carry the highlight, for every pack."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn

from calque_engine.build import build
from calque_engine.pack import load_pack
from calque_engine.style import Style

REPO = Path(__file__).resolve().parents[2]
PACKS = sorted((REPO / "packs").glob("*/pack.yaml"))

CATS = ["North", "South", "East", "West"]
CASES = {
    "line": (
        "trend",
        {
            "categories": CATS,
            "series": [
                {"name": "A", "values": [1, 3, 2, 5]},
                {"name": "B", "values": [2, 2, 4, 3]},
            ],
            "highlight_series": 1,
            "unit": "k",
            "y_title": "Volume",
        },
        XL_CHART_TYPE.LINE,
    ),
    "bar": (
        "quantity",
        {"categories": CATS, "series": [{"name": "A", "values": [4, 7, 3, 5]}], "highlight": 1},
        XL_CHART_TYPE.COLUMN_CLUSTERED,
    ),
    "bar_horizontal": (
        "ranking",
        {"categories": CATS, "series": [{"name": "A", "values": [9, 7, 4, 2]}], "highlight": 0},
        XL_CHART_TYPE.BAR_CLUSTERED,
    ),
    "doughnut": (
        "share",
        {"categories": CATS, "series": [{"name": "A", "values": [40, 30, 20, 10]}], "highlight": 2},
        XL_CHART_TYPE.DOUGHNUT,
    ),
    "scatter": (
        "quantity",
        {
            "points": [
                {"label": "P1", "x": 1, "y": 2},
                {"label": "P2", "x": 3, "y": 5},
                {"label": "P3", "x": 4, "y": 1},
            ],
            "highlight": 1,
            "x_title": "Cost",
            "y_title": "Value",
        },
        XL_CHART_TYPE.XY_SCATTER,
    ),
}


def deck(pack_id: str, chart_type: str) -> dict:
    message_type, params, _ = CASES[chart_type]
    return {
        "pack_id": pack_id,
        "language": "en",
        "title": "Charts",
        "slides": [
            {
                "id": "cover",
                "message": "Charts",
                "message_type": "cover",
                "form": "cover",
                "source": {"kind": "clone", "role": "cover"},
            },
            {
                "id": "chart",
                "message": "North leads",
                "message_type": message_type,
                "form": chart_type,
                "title": "North leads",
                "source": {"kind": "chart", "type": chart_type, "params": params},
            },
        ],
    }


def build_chart(manifest: Path, chart_type: str, tmp_path: Path):
    pack = load_pack(manifest.parent)
    out = tmp_path / f"{chart_type}.pptx"
    build(deck(pack.id, chart_type), pack, out)
    charts = [sh.chart for sh in Presentation(str(out)).slides[1].shapes if sh.has_chart]
    assert len(charts) == 1
    return Style(pack), charts[0]


@pytest.mark.parametrize("chart_type", list(CASES))
@pytest.mark.parametrize("manifest", PACKS, ids=lambda p: p.parent.name)
def test_native_chart_with_highlight(manifest, chart_type, tmp_path):
    st, chart = build_chart(manifest, chart_type, tmp_path)
    _, params, xl = CASES[chart_type]
    assert chart.chart_type == xl
    hi = st.rgb("highlight")
    ser = chart.plots[0].series
    if chart_type == "line":
        assert ser[params["highlight_series"]].format.line.color.rgb == hi
        assert ser[0].format.line.color.rgb == st.rgb("accent")
    elif chart_type == "scatter":
        assert ser[0].points[params["highlight"]].marker.format.fill.fore_color.rgb == hi
    else:
        idx = params["highlight"]
        if chart_type == "bar_horizontal":  # listed bottom-up so the first category reads on top
            idx = len(params["categories"]) - 1 - idx
        assert ser[0].points[idx].format.fill.fore_color.rgb == hi
        if chart_type != "doughnut":
            assert ser[0].format.fill.fore_color.rgb == st.rgb("accent")


@pytest.mark.parametrize("manifest", PACKS, ids=lambda p: p.parent.name)
def test_chart_has_no_title_and_legend_only_for_several_series(manifest, tmp_path):
    for chart_type in ("bar", "line"):
        _, chart = build_chart(manifest, chart_type, tmp_path)
        assert not chart.has_title
        atd = chart._chartSpace.find(qn("c:chart")).find(qn("c:autoTitleDeleted"))
        assert atd is not None and atd.get("val") == "1"
        assert chart.has_legend == (len(chart.plots[0].series) > 1)


def _sourced(pack_id: str, source: str | None) -> dict:
    spec = deck(pack_id, "bar")
    if source:
        spec["slides"][1]["source"]["params"] = {**CASES["bar"][1], "source": source}
    return spec


@pytest.mark.parametrize("manifest", PACKS, ids=lambda p: p.parent.name)
def test_source_line_under_the_chart(manifest, tmp_path):
    """`params.source` is a caption line above the footer line; the chart stops short of it."""
    pack = load_pack(manifest.parent)
    st = Style(pack)
    build(_sourced(pack.id, "Source: CRM, Sept. 2026"), pack, tmp_path / "s.pptx")
    shapes = list(Presentation(str(tmp_path / "s.pptx")).slides[1].shapes)
    (line,) = [sh for sh in shapes if sh.name == "Source"]
    (frame,) = [sh for sh in shapes if sh.has_chart]
    assert line.text_frame.text == "Source: CRM, Sept. 2026"
    run = line.text_frame.paragraphs[0].runs[0]
    assert run.font.color.rgb == st.rgb("muted") and run.font.size.pt == st.size("caption")
    footer = pack.manifest["grid"]["footer_top_in"]
    assert (line.top + line.height) / 914400 <= footer + 1e-6
    assert frame.top + frame.height <= line.top

    build(_sourced(pack.id, None), pack, tmp_path / "n.pptx")
    plain = Presentation(str(tmp_path / "n.pptx")).slides[1].shapes
    assert not [sh for sh in plain if sh.name == "Source"]
    assert next(sh for sh in plain if sh.has_chart).height > frame.height


@pytest.mark.parametrize("manifest", PACKS, ids=lambda p: p.parent.name)
def test_source_goes_in_the_pack_slot(manifest, tmp_path):
    """A content slide mapping a `source` slot gets the text there; with no source the slot goes."""
    pack = load_pack(manifest.parent)
    n = pack.slides_for("content")[0]
    tslide = next(s for s in pack.template_map["slides"] if s["number"] == n)
    if "body" not in tslide.get("slots", {}):
        pytest.skip("the content slide has no spare text shape to stand in for a source slot")
    body = tslide["slots"].pop("body")
    tslide["slots"]["source"] = body  # the body well stands in for a footnote placeholder
    build(_sourced(pack.id, "Source: survey"), pack, tmp_path / "s.pptx")
    shapes = Presentation(str(tmp_path / "s.pptx")).slides[1].shapes
    slot = next(sh for sh in shapes if sh.shape_id == body)
    assert slot.text_frame.text == "Source: survey" and slot.name == "Source"
    build(_sourced(pack.id, None), pack, tmp_path / "n.pptx")
    shapes = Presentation(str(tmp_path / "n.pptx")).slides[1].shapes
    assert body not in {sh.shape_id for sh in shapes}


@pytest.mark.parametrize("manifest", PACKS, ids=lambda p: p.parent.name)
def test_page_numbers_are_slidenum_fields(manifest, tmp_path):
    """Every page number is a slidenum field whose cached text is the slide's position, so it
    stays right after a reorder in PowerPoint."""
    pack = load_pack(manifest.parent)
    report = build(deck(pack.id, "bar"), pack, tmp_path / "d.pptx")
    prs = Presentation(str(tmp_path / "d.pptx"))
    pages = {s["number"]: s.get("page_number") for s in pack.template_map["slides"]}
    numbered = [(pos, n) for pos, n in report.slides.values() if pages.get(n) is not None]
    assert numbered
    for pos, n in numbered:
        shape = next(sh for sh in prs.slides[pos - 1].shapes if sh.shape_id == pages[n])
        fields = shape._element.findall(".//" + qn("a:fld"))
        assert [f.get("type") for f in fields] == ["slidenum"]
        assert shape.text_frame.text.strip() == str(pos)
