"""Generic lint checks, run on every pack in packs/: the template lints clean, and each fault
injected into a cleanly built deck is caught by the right check."""

import io
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from calque_engine.build import build
from calque_engine.lint import lint
from calque_engine.pack import load_pack

REPO = Path(__file__).resolve().parents[2]
PACKS = sorted(p.parent for p in (REPO / "packs").glob("*/pack.yaml"))
POS = 2  # the content slide of the clean deck


def errors(findings):
    return [str(f) for f in findings if f.severity == "ERROR"]


@pytest.fixture(scope="module", params=PACKS, ids=lambda p: p.name)
def clean(request, tmp_path_factory):
    """(pack, path, slide position -> template slide) of a small deck built from the pack."""
    pack = load_pack(request.param)
    slots = {s["number"]: s.get("slots", {}) for s in pack.template_map["slides"]}

    def clone(sid, form, mtype, role, title):
        slot = slots[pack.slides_for(role)[0]].get("title")
        values = {str(slot): title} if slot else {}
        src = {"kind": "clone", "role": role, "values": values}
        return {"id": sid, "message": title, "message_type": mtype, "form": form, "source": src}

    spec = {
        "pack_id": pack.id,
        "language": "en",
        "title": "Review",
        "slides": [
            clone("s1", "cover", "cover", "cover", "Quarterly review"),
            clone("s2", "content", "narrative", "content", "Margins rose in three plants"),
            clone("s3", "closing", "closing", "closing", "Closing"),
        ],
    }
    out = tmp_path_factory.mktemp(pack.id) / "clean.pptx"
    report = build(spec, pack, out)
    return pack, out, dict(report.slides.values())


def test_template_lints_clean(clean):
    pack = clean[0]
    assert errors(lint(pack.template, pack, template=True)) == []


def test_clean_deck_has_no_error(clean):
    pack, path, tmap = clean
    assert errors(lint(path, pack, "en", tmap)) == []


def _box(slide, pack, left, top, width, height, text, size=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tb.text_frame.word_wrap = True
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = pack.font("body")
    if size:
        run.font.size = Pt(size)
    return run


def _off_palette(pack):
    return next(h for h in (f"{i:02X}" * 3 for i in range(1, 255)) if h not in pack.palette())


def f_palette(s, pack, w, h):
    _box(s, pack, 1, 2, 3, 0.4, "colour").font.color.rgb = RGBColor.from_string(_off_palette(pack))


def f_font(s, pack, w, h):
    _box(s, pack, 1, 2, 3, 0.4, "face").font.name = "Nonexistent Sans"


def f_single_use(s, pack, w, h):
    paths = pack.manifest["lint"].get("single_use_colors")
    if not paths:
        pytest.skip("pack declares no single-use colour")
    for left in (1, 4):
        _box(s, pack, left, 2, 1, 0.4, "active").font.color.rgb = RGBColor.from_string(
            pack.color_at(paths[0])
        )


def f_placeholder(s, pack, w, h):
    _box(s, pack, 1, 2, 4, 0.4, pack.manifest["lint"]["placeholders"][0])


def f_page_number(s, pack, w, h):
    _box(s, pack, w * 0.92, h * 0.92, 0.4, 0.3, "9")


def f_off_canvas(s, pack, w, h):
    _box(s, pack, w - 0.5, 2, 2, 0.4, "outside")


def f_footer_band(s, pack, w, h):
    _box(s, pack, 1, pack.manifest["grid"]["footer_top_in"] + 0.1, 2, 0.05, "low")


def f_slop(s, pack, w, h):
    _box(s, pack, 1, 2, 6, 0.6, "In today's fast-paced world, margins rose")


def f_overflow(s, pack, w, h):
    _box(s, pack, 1, 2, 1, 0.3, "A sentence far too long for this narrow box to hold.", 14)


def f_notes(s, pack, w, h):
    s.notes_slide.notes_text_frame.text = "Let's dive in."


FAULTS = {
    "palette": f_palette,
    "font": f_font,
    "single-use": f_single_use,
    "placeholder": f_placeholder,
    "page-number": f_page_number,
    "off-canvas": f_off_canvas,
    "footer-band": f_footer_band,
    "slop": f_slop,
    "overflow": f_overflow,
    "notes": f_notes,
}


@pytest.mark.parametrize("fault", FAULTS)
def test_fault_detected(clean, fault, tmp_path):
    pack, path, tmap = clean
    prs = Presentation(str(path))
    w, h = prs.slide_width / 914400, prs.slide_height / 914400
    FAULTS[fault](prs.slides[POS - 1], pack, w, h)
    bad = tmp_path / f"{fault}.pptx"
    prs.save(str(bad))
    found = lint(bad, pack, "en", tmap)
    before = {(f.check, f.severity) for f in lint(path, pack, "en", tmap) if f.slide == POS}
    new = {(f.check, f.severity) for f in found if f.slide == POS} - before
    expected = ("slop", "NOTE") if fault == "notes" else (fault, None)
    assert any(c == expected[0] and expected[1] in (None, s) for c, s in new), found


def _deck(pack, slides, out):
    spec = {"pack_id": pack.id, "language": "en", "title": "T", "slides": slides}
    return dict(build(spec, pack, out).slides.values())


# signature role -> the message type its form carries
MTYPES = {"cover": "cover", "summary": "summary", "divider": "divider", "closing": "closing"}


def _clone(sid, role, values):
    src = {"kind": "clone", "role": role, "values": values}
    return {"id": sid, "message": sid, "message_type": MTYPES[role], "form": role, "source": src}


def test_rewritten_closing_text_is_linted(clean, tmp_path):
    """Only the closing line the template writes is exempt: the author's copy there is linted,
    and the same with or without a build map."""
    pack = clean[0]
    n = pack.slides_for("closing")[0]
    slots = next(s for s in pack.template_map["slides"] if s["number"] == n).get("slots", {})
    if not slots:
        pytest.skip("closing slide declares no slot")
    values = {str(next(iter(slots.values()))): "In today's fast-paced world, margins rose"}
    tmap = _deck(pack, [_clone("end", "closing", values)], tmp_path / "c.pptx")
    mapped = [str(f) for f in lint(tmp_path / "c.pptx", pack, "en", tmap) if f.check == "slop"]
    unmapped = [str(f) for f in lint(tmp_path / "c.pptx", pack, "en") if f.check == "slop"]
    assert mapped and mapped == unmapped


def _capacity_slot(pack):
    """(role, shape id, capacity) of a slot on a role's first slide that declares a capacity
    and is not a `fit` label."""
    for role, nums in pack.manifest["roles"].items():
        if role not in MTYPES or not nums:
            continue
        s = next(s for s in pack.template_map["slides"] if s["number"] == nums[0])
        shapes = {x["id"]: x for x in s["shapes"]}
        for sid in s.get("slots", {}).values():
            cap = shapes.get(sid, {}).get("capacity")
            if cap and sid not in s.get("fit", []):
                return role, sid, cap
    pytest.skip("pack declares no slot capacity")


def test_slot_capacity_enforced(clean, tmp_path):
    pack = clean[0]
    role, sid, cap = _capacity_slot(pack)
    long = " ".join(["word"] * (cap["chars_per_line"] * (cap["lines"] + 1) // 4 + 2))
    tmap = _deck(pack, [_clone("s", role, {str(sid): long})], tmp_path / "c.pptx")
    for found in (
        lint(tmp_path / "c.pptx", pack, "en", tmap),
        lint(tmp_path / "c.pptx", pack, "en"),
    ):
        assert any(f.check == "capacity" and f.shape_id == sid for f in found), found
    tmap = _deck(pack, [_clone("s", role, {str(sid): "w"})], tmp_path / "d.pptx")
    assert not [f for f in lint(tmp_path / "d.pptx", pack, "en", tmap) if f.check == "capacity"]


def test_char_lines():
    from calque_engine.lint import _char_lines

    assert _char_lines("SUPERCALIFRAGILISTIC WORD", 10) == 3  # a 20-char word overhangs
    assert _char_lines("Document title in two lines maximum", 25) == 2
    assert _char_lines("short", 10) == 1


def test_missing_pack_fonts_are_reported(neutral_pack, tmp_path, capsys):
    """A face with no file in fonts/ renders in its fallback: lint says so, validate-pack too."""
    from calque_engine.__main__ import main

    pack = load_pack(neutral_pack)
    tmap = _deck(pack, [_clone("s", "cover", {})], tmp_path / "c.pptx")
    (fonts,) = [f for f in lint(tmp_path / "c.pptx", pack, "en", tmap) if f.check == "fonts"]
    assert fonts.severity == "WARN" and fonts.slide is None
    assert "rendered with fallback fonts" in fonts.message and pack.font("body") in fonts.message

    manifest = neutral_pack / "pack.yaml"
    data = yaml.safe_load(manifest.read_text())
    same = {r: pack.font(r) for r in ("display", "body")}  # the fallback is the face itself
    data["fonts"] = {"files": ["Face-Regular.ttf"], "fallback": same}
    manifest.write_text(yaml.safe_dump(data))
    found = lint(tmp_path / "c.pptx", load_pack(neutral_pack), "en", tmap)
    assert [f for f in found if f.check == "fonts"] == []
    assert main(["validate-pack", str(neutral_pack)]) == 0
    assert "fonts/ lacks Face-Regular.ttf" in capsys.readouterr().err


RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def test_external_content_and_ole_flagged(clean, tmp_path):
    """DLP: a hyperlink out warns; an OLE object, a remote template and a linked picture fail."""
    from pptx.enum.shapes import PROG_ID

    pack, path, tmap = clean
    prs = Presentation(str(path))
    slide = prs.slides[POS - 1]
    _box(slide, pack, 1, 2, 3, 0.4, "see the source").hyperlink.address = "https://example.org/x"
    link = slide.shapes[-1]
    ole = slide.shapes.add_ole_object(io.BytesIO(b"PK\x05\x06" + b"\0" * 18), PROG_ID.XLSX, 0, 0)
    slide.part.relate_to("https://example.org/beacon.png", f"{RT}/image", is_external=True)
    prs.part.relate_to("https://example.org/remote.potx", f"{RT}/attachedTemplate", True)
    bad = tmp_path / "external.pptx"
    prs.save(str(bad))

    found = [f for f in lint(bad, pack, "en", tmap) if f.check in ("external", "ole")]
    by = {(f.check, f.severity, f.slide): f for f in found}
    assert by[("external", "WARN", POS)].shape_id == link.shape_id
    assert "example.org/x" in by[("external", "WARN", POS)].message
    assert by[("ole", "ERROR", POS)].shape_id == ole.shape_id
    assert any("beacon.png" in f.message and f.severity == "ERROR" for f in found), found
    assert any("remote template" in f.message and f.slide is None for f in found), found
    assert not [f for f in lint(path, pack, "en", tmap) if f.check in ("external", "ole")]
