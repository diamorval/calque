"""Image values are read only from the folders the caller allows, never from anywhere on disk."""

import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from calque_engine.api import call
from calque_engine.build import build
from calque_engine.importer import import_pptx
from calque_engine.pack import load_pack
from calque_engine.slides import read_image

from .test_import_patch import _deck


def _png(path, color):
    Image.new("RGB", (40, 20), color).save(path)
    return path


@pytest.fixture
def imported(neutral_pack, tmp_path):
    """An imported deck whose first slide holds a picture: (pack, deck, base, picture id)."""
    pack = load_pack(neutral_pack)
    src = tmp_path / "src.pptx"
    build(_deck(), pack, src)
    prs = Presentation(str(src))
    pic = prs.slides[0].shapes.add_picture(
        str(_png(tmp_path / "red.png", "red")), Inches(1), Inches(1), Inches(2), Inches(1)
    )
    prs.save(str(src))
    deck, _ = import_pptx(src, pack, tmp_path / "work", language="en")
    return pack, deck, tmp_path / "work" / deck["base"], pic.shape_id


def _with_image(deck, shape_id, ref):
    deck["slides"][0]["source"]["values"] = {str(shape_id): {"image": ref}}
    return deck


def test_read_image_refuses_paths_outside_its_roots(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    _png(root / "ok.png", "blue")
    _png(tmp_path / "secret.png", "green")
    (root / "link.png").symlink_to(tmp_path / "secret.png")
    assert read_image("ok.png", [root])
    for ref in ("/etc/hostname", str(root / "ok.png"), "../secret.png", "sub/../../secret.png"):
        with pytest.raises(ValueError, match="relative path"):
            read_image(ref, [root])
    with pytest.raises(FileNotFoundError):
        read_image("link.png", [root])  # a symlink out of the root
    with pytest.raises(FileNotFoundError):
        read_image("missing.png", [root])


def test_an_image_in_an_allowed_root_replaces_the_picture(imported, tmp_path):
    pack, deck, base, pic = imported
    _png(base.parent / "photo.png", "blue")
    out = tmp_path / "out.pptx"
    build(_with_image(deck, pic, "photo.png"), pack, out, base=base, image_roots=[base.parent])
    shape = next(s for s in Presentation(str(out)).slides[0].shapes if s.shape_id == pic)
    rid = shape._element.find(".//" + qn("a:blip")).get(qn("r:embed"))
    img = Image.open(io.BytesIO(shape.part.related_part(rid).blob)).convert("RGB")
    assert img.getpixel((0, 0)) == (0, 0, 255)


def test_build_call_refuses_images_outside_its_roots(imported, tmp_path):
    pack, deck, base, pic = imported
    for ref in ("/etc/hostname", "../src.pptx"):
        r = call(
            {
                "op": "build",
                "pack": str(pack.dir),
                "deck": _with_image(deck, pic, ref),
                "out": str(tmp_path / "out.pptx"),
                "base": str(base),
                "image_roots": [str(base.parent)],
            }
        )
        assert r["ok"] is False and r["error"] == "ValueError", r
        assert "relative path" in r["message"]
