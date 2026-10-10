"""PowerPoint round trip, on every pack: a built deck edited by hand comes back as a new version of
the same DeckSpec, its drawn slides still drawn (with the client's text and data edits merged),
anything the spec cannot hold kept as an imported clone."""

import copy
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData

from calque_engine import tags
from calque_engine.build import build
from calque_engine.importer import import_deck, import_pptx
from calque_engine.pack import load_pack
from calque_engine.patch import patch
from calque_engine.slides import iter_shapes

REPO = Path(__file__).resolve().parents[2]
PACKS = sorted(p.parent for p in (REPO / "packs").glob("*/pack.yaml"))


def _deck(pack_id: str) -> dict:
    return {
        "pack_id": pack_id,
        "language": "en",
        "title": "Round trip",
        "slides": [
            {
                "id": "cover",
                "message": "Round trip",
                "message_type": "cover",
                "form": "cover",
                "source": {"kind": "clone", "role": "cover", "values": {}},
            },
            {
                "id": "regions",
                "message": "North led growth",
                "message_type": "ranking",
                "form": "bar_horizontal",
                "title": "North led growth",
                "source": {
                    "kind": "chart",
                    "type": "bar_horizontal",
                    "params": {
                        "categories": ["North", "South", "East"],
                        "series": [{"name": "Growth", "values": [18, 7, 5]}],
                        "highlight": 0,
                    },
                },
            },
            {
                "id": "process",
                "message": "Deals close in four steps",
                "message_type": "process",
                "form": "flow",
                "title": "Deals close in four steps",
                "notes": "Say it slowly.",
                "source": {
                    "kind": "diagram",
                    "id": "flow",
                    "params": {"steps": ["Lead", "Qualify", "Proposal", "Close"], "active": 2},
                },
            },
            {
                "id": "plan",
                "message": "Option B is the plan",
                "message_type": "tradeoff",
                "form": "comparison_table",
                "title": "Option B is the plan",
                "source": {
                    "kind": "composition",
                    "id": "comparison_table",
                    "params": {
                        "header": ["Option", "Cost"],
                        "rows": [["Alpha", "40"], ["Beta", "120"]],
                    },
                },
            },
            {
                "id": "end",
                "message": "Thanks",
                "message_type": "closing",
                "form": "closing",
                "source": {"kind": "clone", "role": "closing", "values": {}},
            },
        ],
    }


def _by_id(deck: dict) -> dict:
    return {s["id"]: s for s in deck["slides"]}


def _para(slide, text: str):
    for sh in iter_shapes(slide.shapes):
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if p.text == text:
                    return sh, p
    raise KeyError(text)


@pytest.fixture(params=PACKS, ids=lambda p: p.name)
def built(request, tmp_path):
    pack = load_pack(request.param)
    deck = _deck(pack.id)
    out = tmp_path / "v1.pptx"
    build(deck, pack, out)
    return pack, deck, out


def test_every_slide_is_tagged_and_drawn_ones_carry_their_spec(built):
    _, deck, out = built
    marks = [tags.marks(s) for s in Presentation(str(out)).slides]
    assert [m[0] for m in marks] == [s["id"] for s in deck["slides"]]
    drawn = {m[0]: m[1] for m in marks if m[1]}
    assert set(drawn) == {"regions", "process", "plan"}
    assert drawn["regions"] == _by_id(deck)["regions"]


def test_reimport_merges_text_notes_and_chart_data_into_drawn_slides(built, tmp_path):
    pack, deck, out = built
    prs = Presentation(str(out))
    regions, process, plan = prs.slides[1], prs.slides[2], prs.slides[3]
    for sh in regions.shapes:
        if sh.has_chart:  # the client edits the chart data in PowerPoint (drawn bottom-up)
            data = CategoryChartData()
            data.categories = ["East", "South", "North"]
            data.add_series("Growth", [6, 7, 21])
            sh.chart.replace_data(data)
    _para(regions, "North led growth")[1].runs[0].text = "North leads, by far"
    _para(process, "Qualify")[1].runs[0].text = "Qualify fast"
    process.notes_slide.notes_text_frame.text = "New notes."
    table = next(sh for sh in plan.shapes if sh.has_table)
    table.table.cell(2, 1).text_frame.paragraphs[0].runs[0].text = "110"
    edited = tmp_path / "edited.pptx"
    prs.save(str(edited))

    new, _, report = import_deck(edited, pack, tmp_path / "w", "en", "base-v2.pptx", deck)
    assert report["drawn"] == ["regions", "process", "plan"]
    assert report["imported"] == ["cover", "end"]
    assert report["demoted"] == [] and report["conflicts"] == []
    assert new["title"] == "Round trip" and new["base"] == "base-v2.pptx"
    s = _by_id(new)
    assert [x["id"] for x in new["slides"]] == [x["id"] for x in deck["slides"]]
    assert s["regions"]["title"] == "North leads, by far"
    assert s["regions"]["source"]["params"]["categories"] == ["North", "South", "East"]
    assert s["regions"]["source"]["params"]["series"][0]["values"] == [21, 7, 6]
    assert s["process"]["source"]["params"]["steps"][1] == "Qualify fast"
    assert s["process"]["notes"] == "New notes."
    assert s["plan"]["source"]["params"]["rows"][1] == ["Beta", "110"]
    assert s["cover"]["source"] == {"kind": "clone", "from": "base", "slide": 1, "values": {}}

    # the chart is still a chart: set_params works, and the deck builds on its new base
    patched, _ = patch(
        new, [{"op": "set_params", "slide": "regions", "params": {"highlight": 1}}], pack
    )
    v2 = tmp_path / "v2.pptx"
    r = build(patched, pack, v2, base=tmp_path / "w" / "base-v2.pptx")
    assert list(r.slides) == [x["id"] for x in deck["slides"]]
    rebuilt = Presentation(str(v2))
    assert len(rebuilt.slides) == 5
    _para(rebuilt.slides[2], "Qualify fast")
    assert rebuilt.slides[2].notes_slide.notes_text_frame.text == "New notes."

    # and the rebuilt file round-trips again unchanged
    again, _, rep2 = import_deck(v2, pack, tmp_path / "w", "en", "base-v3.pptx", patched)
    assert rep2["drawn"] == ["regions", "process", "plan"] and rep2["demoted"] == []
    assert _by_id(again)["regions"] == _by_id(patched)["regions"]


def test_changes_the_spec_cannot_hold_keep_the_slide_as_a_clone(built, tmp_path):
    pack, deck, out = built
    prs = Presentation(str(out))
    table = next(sh for sh in prs.slides[3].shapes if sh.has_table)
    table.left += 300000  # moved by hand
    sh, _ = _para(prs.slides[2], "Lead")
    sh.text_frame.add_paragraph().text = "an extra line"  # text the flow cannot hold
    edited = tmp_path / "edited.pptx"
    prs.save(str(edited))

    new, _, report = import_deck(edited, pack, tmp_path / "w", "en", "b.pptx", deck)
    assert report["drawn"] == ["regions"]
    why = {d["slide"]: d["reason"] for d in report["demoted"]}
    assert set(why) == {"process", "plan"}
    assert "moved" in why["plan"] and "paragraphs" in why["process"]
    s = _by_id(new)
    # still anchored on the same ids, so comments and history follow
    assert s["plan"]["source"] == {"kind": "clone", "from": "base", "slide": 4, "values": {}}
    assert s["plan"]["form"] == "imported"
    out2 = tmp_path / "v2.pptx"
    build(new, pack, out2, base=tmp_path / "w" / "b.pptx")
    assert Presentation(str(out2)).slides[3].shapes  # kept exactly as the client left it


def test_duplicated_and_foreign_slides_get_new_ids(built, tmp_path):
    from calque_engine.slides import duplicate_slide

    pack, deck, out = built
    prs = Presentation(str(out))
    duplicate_slide(prs, prs.slides[1])  # duplicated in PowerPoint: the tags come along
    prs.save(str(tmp_path / "dup.pptx"))
    previous = copy.deepcopy(deck)
    previous["slides"].append({**previous["slides"][0], "id": "s6"})  # an id already taken
    new, _, report = import_deck(
        tmp_path / "dup.pptx", pack, tmp_path / "w", "en", "b.pptx", previous
    )
    ids = [s["id"] for s in new["slides"]]
    assert ids[:5] == [s["id"] for s in deck["slides"]]
    assert ids[5] == "s6-2" and ids[5] in report["imported"]


def test_conflict_when_calque_changed_the_slide_after_export(built, tmp_path):
    pack, deck, out = built
    head = copy.deepcopy(deck)
    _by_id(head)["regions"]["title"] = "Changed in Calque since"
    new, _, report = import_deck(out, pack, tmp_path / "w", "en", "b.pptx", head)
    assert [c["slide"] for c in report["conflicts"]] == ["regions"]
    assert _by_id(new)["regions"]["title"] == "North led growth"  # the file wins


def test_drawn_and_template_slides_can_be_added_to_an_imported_deck(built, tmp_path):
    pack, deck, _ = built
    # a deck that never went through the engine: the pack template itself, no tags
    base_deck, _ = import_pptx(pack.template, pack, tmp_path / "w", "en")
    assert all(s["form"] == "imported" for s in base_deck["slides"])
    chart, cover = _by_id(deck)["regions"], _by_id(deck)["cover"]
    out_deck, _ = patch(
        base_deck,
        [
            {"op": "insert_slide", "at": 0, "slide": cover},
            {"op": "insert_slide", "at": 1, "slide": chart},
        ],
        pack,
    )
    out = tmp_path / "out.pptx"
    r = build(out_deck, pack, out, base=tmp_path / "w" / "base.pptx")
    prs = Presentation(str(out))
    assert len(prs.slides) == len(base_deck["slides"]) + 2
    assert any(sh.has_chart for sh in prs.slides[1].shapes)
    assert r.slides["regions"][0] == 2


def test_a_restyled_drawn_slide_is_kept_as_the_client_left_it(built, tmp_path):
    pack, deck, out = built
    prs = Presentation(str(out))
    _para(prs.slides[2], "Close")[1].runs[0].font.italic = True
    prs.save(str(tmp_path / "edited.pptx"))
    _, _, report = import_deck(tmp_path / "edited.pptx", pack, tmp_path / "w", "en", "b.pptx", deck)
    assert [d["slide"] for d in report["demoted"]] == ["process"]
    assert "restyled" in report["demoted"][0]["reason"]
