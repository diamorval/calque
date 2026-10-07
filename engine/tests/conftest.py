"""A neutral pack built on python-pptx's default template: proves the engine needs no brand."""

import json
from pathlib import Path

import pytest
import yaml
from pptx import Presentation

from calque_engine.extract import draft_tokens, extract

REPO = Path(__file__).resolve().parents[2]


def make_template(path: Path) -> Path:
    prs = Presentation()
    for layout_idx, title in ((0, "Deck title"), (1, "Content title"), (0, "Closing")):
        s = prs.slides.add_slide(prs.slide_layouts[layout_idx])
        s.shapes.title.text = title
    prs.save(path)
    return path


@pytest.fixture
def neutral_pack(tmp_path: Path) -> Path:
    d = tmp_path / "neutral"
    d.mkdir()
    make_template(d / "template.pptx")
    tmap = extract(d / "template.pptx")
    (d / "template-map.yaml").write_text(yaml.safe_dump(tmap, sort_keys=False))
    (d / "tokens.json").write_text(json.dumps(draft_tokens(tmap)))
    (d / "pack.yaml").write_text(
        yaml.safe_dump(
            {
                "id": "neutral",
                "name": "Neutral",
                "version": "0.1.0",
                "default_language": "en",
                "missing_value": {"en": "[TO COMPLETE]"},
                "roles": {"cover": [1], "content": [2], "closing": [3]},
                "never_clone": [],
                "grid": {
                    "margin_in": 0.5,
                    "columns": {"3": [0.5, 3.6, 6.7]},
                    "title": {"left_in": 0.5, "top_in": 0.3, "width_in": 9, "height_in": 1.2},
                    "body_top_in": 1.8,
                    "footer_top_in": 7.0,
                },
                "fonts": {"fallback": {"display": "serif", "body": "sans-serif"}},
                "lint": {},
            }
        )
    )
    return d
