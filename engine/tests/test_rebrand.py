"""Re-brand a deck from one pack to another, for every ordered pair of packs in packs/."""

import itertools

import pytest
from pptx import Presentation

from calque_engine.build import build
from calque_engine.importer import import_pptx
from calque_engine.pack import load_pack
from calque_engine.rebrand import RebrandError, rebrand

from .test_roundtrip import PACKS, _by_id, _deck, _para

PAIRS = list(itertools.permutations(PACKS, 2))


def _slot(pack, role: str, name: str) -> int:
    n = pack.slides_for(role)[0]
    return next(s for s in pack.template_map["slides"] if s["number"] == n)["slots"][name]


@pytest.mark.skipif(not PAIRS, reason="needs two packs")
@pytest.mark.parametrize(("a", "b"), PAIRS, ids=lambda p: p.name)
def test_rebrand_redraws_drawn_slides_and_moves_clones_by_role(a, b, tmp_path):
    old, new = load_pack(a), load_pack(b)
    deck = _deck(old.id)
    cover = _by_id(deck)["cover"]["source"]
    cover["values"] = {str(_slot(old, "cover", "title")): "Round trip", "999": "not a slot"}

    out, report = rebrand(deck, old, new)
    assert out["pack_id"] == new.id
    assert report["redrawn"] == ["regions", "process", "plan"]
    assert report["moved"] == ["cover", "end"]
    assert report["unmapped"] == [{"slide": "cover", "shape_id": 999, "slot": None}]
    assert _by_id(out)["cover"]["source"]["values"] == {
        str(_slot(new, "cover", "title")): "Round trip"
    }
    assert _by_id(out)["regions"] == _by_id(deck)["regions"]
    build(out, new, tmp_path / "out.pptx")
    prs = Presentation(str(tmp_path / "out.pptx"))
    assert len(prs.slides) == 5
    _para(prs.slides[0], "Round trip")


@pytest.mark.skipif(not PAIRS, reason="needs two packs")
@pytest.mark.parametrize(("a", "b"), PAIRS[:1], ids=lambda p: p.name)
def test_imported_slides_are_not_rebrandable(a, b, tmp_path):
    old, new = load_pack(a), load_pack(b)
    imported, _ = import_pptx(old.template, old, tmp_path / "w", "en")
    deck = {**imported, "slides": [*imported["slides"][:1], _by_id(_deck(old.id))["regions"]]}
    with pytest.raises(RebrandError, match="drop"):
        rebrand(deck, old, new)
    out, report = rebrand(deck, old, new, drop=True)
    assert report["dropped"] == [imported["slides"][0]["id"]]
    assert [s["id"] for s in out["slides"]] == ["regions"] and "base" not in out
