"""Plain text of a document the user attached, for the model: minimal, no layout.

txt/md/csv are read as is; pptx with python-pptx; docx and xlsx straight from their XML parts.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
X = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
PLAIN = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml"}


def _docx(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    paras = ("".join(t.text or "" for t in p.iter(f"{W}t")) for p in root.iter(f"{W}p"))
    return "\n".join(p for p in paras if p.strip())


def _xlsx(path: Path) -> str:
    out: list[str] = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            sst = etree.fromstring(z.read("xl/sharedStrings.xml"))
            shared = ["".join(t.text or "" for t in si.iter(f"{X}t")) for si in sst.iter(f"{X}si")]
        sheets = sorted(
            (n for n in names if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")),
            key=lambda n: int("".join(c for c in n if c.isdigit()) or 0),
        )
        for i, name in enumerate(sheets, 1):
            out.append(f"# Sheet {i}")
            for row in etree.fromstring(z.read(name)).iter(f"{X}row"):
                cells = []
                for c in row.iter(f"{X}c"):
                    v = c.find(f"{X}v")
                    if c.get("t") == "s" and v is not None:
                        cells.append(shared[int(v.text)])
                    elif c.get("t") == "inlineStr":
                        cells.append("".join(t.text or "" for t in c.iter(f"{X}t")))
                    else:
                        cells.append(v.text if v is not None and v.text else "")
                if any(cells):
                    out.append(",".join(cells))
    return "\n".join(out)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    out: list[str] = []
    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        out.append(f"# Slide {n}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


def read_text(path: str | Path, name: str) -> str:
    """The text of `path`, its type taken from the file `name`'s extension."""
    ext = Path(name).suffix.lower()
    p = Path(path)
    if ext in PLAIN:
        return p.read_text(encoding="utf-8", errors="replace")
    if ext == ".docx":
        return _docx(p)
    if ext == ".xlsx":
        return _xlsx(p)
    if ext == ".pptx":
        return _pptx(p)
    raise NotImplementedError(f"reading {ext or 'this'} files is not supported yet")
