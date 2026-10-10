"""The engine's one interface for the server: JSON in, JSON out (`python -m calque_engine call`).

Request: {"op": <name>, ...args}. Response: {"ok": true, ...result} or
{"ok": false, "error": <kind>, "message": str, "issues": [...]}. Packs are passed as directories.
"""

from __future__ import annotations

from dataclasses import asdict
from typing import Any

from .deckspec import DeckSpecError, Issue, validate
from .pack import PackError, load_pack
from .patch import PatchError


def _issues(issues: list[Issue]) -> list[dict[str, Any]]:
    return [asdict(i) for i in issues]


def op_validate_pack(pack: str) -> dict[str, Any]:
    p = load_pack(pack)
    return {"id": p.id, "version": p.manifest["version"]}


def op_validate(pack: str, deck: dict[str, Any]) -> dict[str, Any]:
    _, warnings = validate(deck, load_pack(pack))
    return {"warnings": _issues(warnings)}


def op_build(
    pack: str,
    deck: dict[str, Any],
    out: str,
    base: str | None = None,
    image_roots: list[str] | None = None,
    author: str | None = None,
) -> dict[str, Any]:
    """`image_roots`: the folders image values may be read from (see `build`)."""
    from .build import build

    r = build(deck, load_pack(pack), out, base=base, image_roots=image_roots, author=author)
    return {
        "path": str(r.path),
        "warnings": _issues(r.warnings),
        "holes": [asdict(h) for h in r.holes],
        "slides": {k: {"position": v[0], "source": v[1]} for k, v in r.slides.items()},
    }


def op_patch(pack: str, deck: dict[str, Any], ops: list[dict[str, Any]]) -> dict[str, Any]:
    from .patch import patch

    out, warnings = patch(deck, ops, load_pack(pack))
    return {"deck": out, "warnings": _issues(warnings)}


def op_import(
    pack: str,
    pptx: str,
    dest: str,
    language: str,
    base_id: str = "base.pptx",
    previous: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """`previous`: the DeckSpec this file is re-imported into (its slides are recognised by the
    tags the engine wrote). `report`: slides kept drawn, imported, demoted (why), conflicts."""
    from .importer import import_deck

    deck, tmap, report = import_deck(pptx, load_pack(pack), dest, language, base_id, previous)
    return {"deck": deck, "template_map": tmap, "report": report}


def op_graft(
    pack: str, out: str, sources: list[dict[str, Any]], base: str | None = None
) -> dict[str, Any]:
    """Copy slides of other decks' files into a deck's base, for copy_slides: `out` = `base`
    (default: the pack template) + each `{"pptx": path, "slide": n}` appended, its layout matched
    by name. Returns the appended slides' numbers in `out`."""
    from pptx import Presentation

    from .slides import duplicate_slide

    prs = Presentation(base or str(load_pack(pack).template))
    start = len(prs.slides)
    opened: dict[str, Any] = {}
    for src in sources:
        if src["pptx"] not in opened:
            opened[src["pptx"]] = Presentation(src["pptx"])
        sp, n = opened[src["pptx"]], src["slide"]
        if not 1 <= n <= len(sp.slides):
            raise ValueError(f"slide {n} not in its deck's file (1..{len(sp.slides)})")
        duplicate_slide(prs, sp.slides[n - 1])
    prs.save(out)
    return {"path": out, "slides": list(range(start + 1, start + 1 + len(sources)))}


def op_extract(
    template: str,
    pack_id: str = "draft",
    name: str = "Draft",
    tokens: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Drafts for a new pack: template map (with title slots on role slides), tokens, manifest.
    A .potx is rewritten in place as a .pptx. `tokens`: the company's own tokens.json, used
    instead of the draft. `review`: the resolved colours and fonts, for a person to check."""
    from . import tokens as tk
    from .extract import as_presentation, draft_manifest, draft_tokens, extract, used_styles

    as_presentation(template)
    tmap = extract(template)
    manifest = draft_manifest(tmap, pack_id, name, used_styles(template))
    tokens = tokens or draft_tokens(tmap)
    values, flat = tk.resolve(tokens), tk.flatten(tokens)
    review = {
        kind: {p: fn(values[p]) for p, t in flat.items() if t.get("$type") == typ}
        for kind, typ, fn in (("colors", "color", tk.hex6), ("fonts", "fontFamily", tk.family))
    }
    return {"template_map": tmap, "tokens": tokens, "manifest": manifest, "review": review}


def op_lint(
    pack: str,
    pptx: str,
    language: str | None = None,
    sources: dict[str, int] | None = None,
    template: bool = False,
) -> dict[str, Any]:
    """`sources`: output position -> template slide number, as returned by build.
    `template`: lint the pack's own template (placeholders and slop are its job)."""
    from .fix import SAFE
    from .lint import lint

    p = load_pack(pack)
    tmap = {int(k): v for k, v in sources.items()} if sources else None
    findings = lint(pptx, p, language=language, template_map=tmap, template=template)
    return {"findings": [asdict(f) for f in findings], "safe_checks": list(SAFE)}


def op_fix(pack: str, pptx: str, out: str) -> dict[str, Any]:
    """Apply the safe lint fixes (`fix.SAFE`) to `pptx`, written to `out`."""
    from .fix import SAFE, fix

    return {"path": out, "applied": fix(pptx, load_pack(pack), out), "safe": list(SAFE)}


def op_render(
    pack: str,
    pptx: str,
    out_dir: str,
    only: list[int] | None = None,
    dpi: int = 96,
    sources: dict[str, int] | None = None,
) -> dict[str, Any]:
    """`sources`: output position -> template slide number (from build), for shape roles."""
    from .render import render

    r = render(
        pptx,
        out_dir,
        load_pack(pack),
        dpi=dpi,
        only=set(only) if only else None,
        sources={int(k): v for k, v in sources.items()} if sources else None,
    )
    return asdict(r) if hasattr(r, "__dataclass_fields__") else dict(r)


def op_text(path: str, name: str, limit: int = 100_000) -> dict[str, Any]:
    """Plain text of an attached document (txt, md, csv, docx, xlsx, pptx), cut at `limit`."""
    from .text import read_text

    text = read_text(path, name)
    return {"text": text[:limit], "truncated": len(text) > limit}


OPS = {
    "validate_pack": op_validate_pack,
    "validate": op_validate,
    "build": op_build,
    "patch": op_patch,
    "import": op_import,
    "extract": op_extract,
    "graft": op_graft,
    "lint": op_lint,
    "fix": op_fix,
    "render": op_render,
    "text": op_text,
}


def call(request: dict[str, Any]) -> dict[str, Any]:
    args = dict(request)
    op = args.pop("op", None)
    fn = OPS.get(op)
    if fn is None:
        return {
            "ok": False,
            "error": "unknown_op",
            "message": f"unknown op {op!r}",
            "ops": sorted(OPS),
        }
    try:
        return {"ok": True, **fn(**args)}
    except DeckSpecError as e:
        return {
            "ok": False,
            "error": "invalid_deck",
            "message": str(e),
            "issues": _issues(e.issues),
        }
    except PackError as e:
        return {"ok": False, "error": "invalid_pack", "message": str(e), "issues": e.problems}
    except PatchError as e:
        return {"ok": False, "error": "invalid_patch", "message": str(e)}
    except TypeError as e:
        return {"ok": False, "error": "bad_request", "message": str(e)}
    except (KeyError, ValueError, FileNotFoundError, NotImplementedError) as e:
        return {"ok": False, "error": type(e).__name__, "message": str(e)}
