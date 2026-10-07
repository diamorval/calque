"""PPTX -> one PNG per slide (headless LibreOffice + pdftoppm) and a shape map, incrementally.

Each slide is keyed by a hash of its XML, every part it reaches (layout, master, theme, images,
charts), its position (slide-number fields), the canvas, the DPI and the pack fonts. Only slides
whose key changed are converted, via the PDF export `PageRange` on the unmodified deck, so page
numbers and inherited layouts render exactly as in the full deck.

Fonts: every call gets its own throwaway LibreOffice profile (no lock contention, nothing written
to the user's system) and the pack's `fonts/` are copied into `<profile>/user/fonts`, which
LibreOffice loads at start on macOS and Linux. A role family the pack ships no file for is
rewritten to its `fonts.fallback` family in a temporary copy of the deck.
"""

from __future__ import annotations

import fcntl
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from .extract import shape_kind
from .pack import Pack
from .style import Style

CACHE_VERSION = 1
INDEX = "render-cache.json"
SOFFICE_TIMEOUT = 300
PDFTOPPM_TIMEOUT = 120
SKIP_RELS = {RT.NOTES_SLIDE, RT.SLIDE}  # notes don't render; a link to a slide isn't its content


class RenderError(RuntimeError):
    pass


@dataclass
class RenderResult:
    pngs: list[Path]  # slide-<n>.png, in deck order
    shapes_path: Path
    slides: list[dict[str, Any]]  # the shape map, as written to shapes.json
    rendered: list[int] = field(default_factory=list)  # 1-based slide numbers converted now
    cached: list[int] = field(default_factory=list)  # 1-based slide numbers reused


def _tool(name: str, mac_app: str | None = None) -> str:
    found = shutil.which(name)
    if found:
        return found
    if mac_app and Path(mac_app).is_file():
        return mac_app
    raise RenderError(f"{name} not found: install LibreOffice and poppler to render slides")


def _fonts_digest(pack: Pack) -> str:
    h = hashlib.sha256(json.dumps(pack.manifest.get("fonts", {}), sort_keys=True).encode())
    d = pack.dir / "fonts"
    for f in sorted(d.iterdir()) if d.is_dir() else []:
        h.update(f.name.encode())
        h.update(f.read_bytes())
    return h.hexdigest()


def font_swaps(pack: Pack) -> dict[str, str]:
    """Role family -> fallback family, for each role family the pack ships no font file for."""
    style = Style(pack)
    out = {}
    for role, fallback in pack.manifest["fonts"]["fallback"].items():
        try:
            family = pack.font(role)
        except KeyError:
            continue
        if family != fallback and style.font_file(family) is None:
            out[family] = fallback
    return out


def _swap_fonts(src: Path, dst: Path, swaps: dict[str, str]) -> None:
    pat = re.compile(rb'typeface="(' + b"|".join(re.escape(f.encode()) for f in swaps) + rb')"')
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item)
            if item.filename.endswith(".xml"):
                data = pat.sub(lambda m: b'typeface="%s"' % swaps[m[1].decode()].encode(), data)
            zout.writestr(item, data)


def _slide_keys(prs, salt: str) -> list[str]:
    seen: dict[int, str] = {}

    def part_digest(part) -> str:
        if id(part) not in seen:
            seen[id(part)] = hashlib.sha256(part.blob).hexdigest()
        return seen[id(part)]

    keys = []
    for pos, slide in enumerate(prs.slides, start=1):
        digests, stack, visited = [], [slide.part], set()
        while stack:
            part = stack.pop()
            if id(part) in visited:
                continue
            visited.add(id(part))
            digests.append(f"{part.partname}:{part_digest(part)}")
            for rel in part.rels.values():
                if not rel.is_external and rel.reltype not in SKIP_RELS:
                    stack.append(rel.target_part)
        h = hashlib.sha256(f"{salt}|{pos}".encode())
        for d in sorted(digests):
            h.update(d.encode())
        keys.append(h.hexdigest())
    return keys


# --- shape map -------------------------------------------------------------------------------


def _xy(el, tag: str, a: str, b: str) -> tuple[int, int]:
    e = el.find(qn(tag))
    return (int(e.get(a)), int(e.get(b))) if e is not None else (0, 0)


def _walk(shapes, tf, out: list):
    """(shape, slide-EMU bbox) for every shape, groups included, through group transforms."""
    ox, oy, sx, sy = tf
    for sh in shapes:
        if sh.left is not None and sh.width is not None:
            bbox = (ox + sh.left * sx, oy + sh.top * sy, sh.width * sx, sh.height * sy)
            out.append((sh, bbox))
        if sh.element.tag == qn("p:grpSp"):
            xfrm = sh.element.grpSpPr.find(qn("a:xfrm"))
            if xfrm is None:
                _walk(sh.shapes, tf, out)
                continue
            off, ext = _xy(xfrm, "a:off", "x", "y"), _xy(xfrm, "a:ext", "cx", "cy")
            choff, chext = _xy(xfrm, "a:chOff", "x", "y"), _xy(xfrm, "a:chExt", "cx", "cy")
            kx = ext[0] / chext[0] if chext[0] else 1.0
            ky = ext[1] / chext[1] if chext[1] else 1.0
            child_tf = (
                ox + sx * (off[0] - choff[0] * kx),
                oy + sy * (off[1] - choff[1] * ky),
                sx * kx,
                sy * ky,
            )
            _walk(sh.shapes, child_tf, out)
    return out


def _clip(bbox, scale_x: float, scale_y: float, w: int, h: int) -> list[int]:
    x0, y0 = bbox[0] * scale_x, bbox[1] * scale_y
    x1, y1 = x0 + bbox[2] * scale_x, y0 + bbox[3] * scale_y
    x0, x1 = sorted((min(max(x0, 0), w), min(max(x1, 0), w)))
    y0, y1 = sorted((min(max(y0, 0), h), min(max(y1, 0), h)))
    return [round(x0), round(y0), round(x1) - round(x0), round(y1) - round(y0)]


def _ids(shapes):
    for sh in shapes:
        yield sh["id"]
        yield from _ids(sh.get("children", []))


def _shape_map(prs, pngs, sources, tmap) -> list[dict[str, Any]]:
    by_n = {s["number"]: s for s in tmap.get("slides", [])}
    out = []
    for pos, (slide, png) in enumerate(zip(prs.slides, pngs, strict=True), start=1):
        with Image.open(png) as im:
            w, h = im.size
        src = by_n.get(sources.get(pos)) if sources else None
        slot_of = {v: k for k, v in (src or {}).get("slots", {}).items()}
        template_ids = set(_ids(src["shapes"])) if src else set()
        shapes = []
        # ponytail: rotation ignored, the bbox is the unrotated frame
        for sh, bbox in _walk(slide.shapes, (0, 0, 1.0, 1.0), []):
            sid = sh.shape_id
            role = slot_of.get(sid) or ("template" if sid in template_ids else "drawn")
            shapes.append(
                {
                    "shape_id": sid,
                    "bbox_px": _clip(bbox, w / prs.slide_width, h / prs.slide_height, w, h),
                    "kind": shape_kind(sh),
                    "role": role,
                }
            )
        out.append({"number": pos, "width_px": w, "height_px": h, "shapes": shapes})
    return out


# --- conversion ------------------------------------------------------------------------------


def _convert(deck: Path, pack: Pack, pages: list[int], dpi: int, work: Path) -> list[Path]:
    """Render `pages` (1-based) of `deck` to PNGs, in that order."""
    soffice = _tool("soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice")
    pdftoppm = _tool("pdftoppm")
    profile = work / "profile"
    (profile / "user").mkdir(parents=True)
    fonts = pack.dir / "fonts"
    if fonts.is_dir():
        shutil.copytree(fonts, profile / "user" / "fonts")
    swaps = font_swaps(pack)
    if swaps:
        swapped = work / "deck.pptx"
        _swap_fonts(deck, swapped, swaps)
        deck = swapped
    opts = {
        "PageRange": {"type": "string", "value": ",".join(map(str, pages))},
        "ExportHiddenSlides": {"type": "boolean", "value": "true"},
    }
    pdf_dir = work / "pdf"
    cmd = [
        soffice,
        f"-env:UserInstallation={profile.as_uri()}",
        "--headless",
        "--norestore",
        "--convert-to",
        "pdf:impress_pdf_Export:" + json.dumps(opts),
        "--outdir",
        str(pdf_dir),
        str(deck),
    ]
    run = subprocess.run(cmd, capture_output=True, text=True, timeout=SOFFICE_TIMEOUT)
    pdf = pdf_dir / (deck.stem + ".pdf")
    if not pdf.is_file():
        raise RenderError(f"LibreOffice produced no PDF for {deck}: {run.stderr or run.stdout}")
    subprocess.run(
        [pdftoppm, "-png", "-r", str(dpi), str(pdf), str(work / "page")],
        check=True,
        capture_output=True,
        timeout=PDFTOPPM_TIMEOUT,
    )
    out = sorted(work.glob("page-*.png"), key=lambda p: int(p.stem.rsplit("-", 1)[1]))
    if len(out) != len(pages):
        raise RenderError(f"expected {len(pages)} pages from {deck}, got {len(out)}")
    return out


def render(
    pptx_path: str | Path,
    out_dir: str | Path,
    pack: Pack,
    dpi: int = 96,
    only: set[int] | None = None,
    template_map: dict[str, Any] | None = None,
    sources: dict[int, int] | None = None,
) -> RenderResult:
    """Render `slide-<n>.png` per slide and `shapes.json` into `out_dir`, reusing unchanged slides.

    `only` forces those slides (1-based) to be re-rendered and leaves the others as they are
    (a slide with no PNG yet is rendered anyway).
    `sources` maps slide position -> template slide number (`BuildReport.slides`), so shapes on a
    reviewed slot of `template_map` (default: the pack's) get the slot name as their role.
    """
    deck = Path(pptx_path)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(deck))
    n = len(prs.slides)
    salt = f"{CACHE_VERSION}|{dpi}|{prs.slide_width}x{prs.slide_height}|{_fonts_digest(pack)}"
    keys = _slide_keys(prs, salt)

    with open(out / ".lock", "w") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)  # two renders into one out_dir take turns
        index_path = out / INDEX
        try:
            index = json.loads(index_path.read_text())
        except (OSError, ValueError):
            index = {}
        cached_keys: dict[str, str] = index.get("slides", {})

        def png(i: int) -> Path:
            return out / f"slide-{i}.png"

        if only is not None:
            bad = sorted(i for i in only if not 1 <= i <= n)
            if bad:
                raise ValueError(f"slides {bad} not in deck (1..{n})")
            todo = sorted(only | {i for i in range(1, n + 1) if not png(i).is_file()})
        else:
            todo = [
                i
                for i in range(1, n + 1)
                if cached_keys.get(str(i)) != keys[i - 1] or not png(i).is_file()
            ]
        if todo:
            with tempfile.TemporaryDirectory(prefix="calque-render-") as tmp:
                for i, page in zip(todo, _convert(deck, pack, todo, dpi, Path(tmp)), strict=True):
                    shutil.move(page, png(i))
                    cached_keys[str(i)] = keys[i - 1]
        for i in [int(k) for k in cached_keys if int(k) > n]:
            png(i).unlink(missing_ok=True)
            del cached_keys[str(i)]
        index_path.write_text(json.dumps({"version": CACHE_VERSION, "slides": cached_keys}))

        pngs = [png(i) for i in range(1, n + 1)]
        slides = _shape_map(prs, pngs, sources, template_map or pack.template_map)
        shapes_path = out / "shapes.json"
        shapes_path.write_text(json.dumps({"slides": slides}, indent=1))
    done = set(todo)
    return RenderResult(
        pngs=pngs,
        shapes_path=shapes_path,
        slides=slides,
        rendered=todo,
        cached=[i for i in range(1, n + 1) if i not in done],
    )
