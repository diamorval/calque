"""import_pptx: turn an existing deck into a DeckSpec that edits it in place (`base`).

A deck the engine built carries tags on its slides (`tags.py`): the slide id, and for a drawn
slide the spec it was drawn from. When such a deck comes back from PowerPoint:

- a tagged slide keeps its id, so comments and history stay anchored on it;
- a drawn slide (chart, diagram, composition) stays drawn when the client only changed its text,
  its notes or its chart data: those edits are merged into the spec, so `set_params` keeps working.
  The slide is redrawn from the spec alone on the pack, and compared shape by shape (shape ids,
  geometry, text, chart data, pictures) with what came back;
- anything else (a shape added, removed, moved or resized, a text the spec cannot hold, a slide
  without tags) becomes an `imported` clone of the file, kept exactly as the client left it.
"""

from __future__ import annotations

import copy
import hashlib
import re
import shutil
import tempfile
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from . import tags
from .deckspec import DeckSpecError, validate
from .extract import extract
from .pack import Pack
from .slides import iter_shapes

DRAWN = ("chart", "diagram", "composition")
ID_RE = re.compile(r"^[A-Za-z0-9_-]+$")
TOLERANCE_EMU = 12700  # 1 pt: below what a hand move in PowerPoint produces


def _title(slide) -> str:
    """The slide's title placeholder, else its topmost text shape."""
    if slide.shapes.title is not None and slide.shapes.title.text_frame.text.strip():
        return slide.shapes.title.text_frame.text.strip()
    texts = [
        sh
        for sh in iter_shapes(slide.shapes)
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.top is not None
    ]
    if not texts:
        return ""
    return min(texts, key=lambda sh: (sh.top, sh.left)).text_frame.text.strip()


def import_pptx(
    pptx: str | Path, pack: Pack, dest: str | Path, language: str, base_id: str = "base.pptx"
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Copy `pptx` to `dest/base_id` and return (DeckSpec, template map of the imported deck)."""
    deck, tmap, _ = import_deck(pptx, pack, dest, language, base_id)
    return deck, tmap


def import_deck(
    pptx: str | Path,
    pack: Pack,
    dest: str | Path,
    language: str,
    base_id: str = "base.pptx",
    previous: dict[str, Any] | None = None,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    """Copy `pptx` to `dest/base_id`; return (DeckSpec, template map of the file, report).

    `previous` is the DeckSpec the file is re-imported into (a new version of that deck): its
    title is kept and its ids are not reused for new slides. The report lists the slides kept
    `drawn`, those `imported` as clones, the drawn ones `demoted` to clones (with the reason) and
    `conflicts`: drawn slides changed in Calque after this file was exported (the file wins).
    Unrecognised slides become clones from the base with no values: the deck rebuilds identical.
    """
    dest = Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    base = dest / base_id
    shutil.copyfile(pptx, base)
    prs = Presentation(str(base))
    prev = {s["id"]: s for s in (previous or {}).get("slides", [])}

    entries: list[dict[str, Any]] = []
    seen: set[str] = set()
    for n, slide in enumerate(prs.slides, start=1):
        sid, spec = tags.marks(slide)
        if sid is not None and (sid in seen or not ID_RE.match(sid)):
            sid, spec = None, None  # a slide duplicated in PowerPoint: the first one keeps the id
        if spec is not None and (
            spec.get("id") != sid or (spec.get("source") or {}).get("kind") not in DRAWN
        ):
            spec = None
        if sid is not None:
            seen.add(sid)
        entries.append({"n": n, "slide": slide, "id": sid, "spec": spec})
    taken = seen | set(prev)
    for e in entries:
        if e["id"] is None:
            sid, k = f"s{e['n']}", 1
            while sid in taken:
                k += 1
                sid = f"s{e['n']}-{k}"
            taken.add(sid)
            e["id"] = sid

    report: dict[str, Any] = {"drawn": [], "imported": [], "demoted": [], "conflicts": []}
    with tempfile.TemporaryDirectory() as tmp:
        refs = _references([e["spec"] for e in entries if e["spec"]], pack, language, Path(tmp))
        out_slides: list[dict[str, Any]] = []
        slots = _slots(pack) if refs else {}
        for e in entries:
            spec = e["spec"]
            if spec is not None:
                ref = refs.get(e["id"])
                merged, why = (
                    (None, ref) if isinstance(ref, str) else _merge(spec, ref, e["slide"], slots)
                )
                if merged is not None:
                    e["drawn"] = merged
                    out_slides.append(merged)
                    continue
                report["demoted"].append({"slide": e["id"], "reason": why})
            out_slides.append(_clone(e, prev))

        deck = {
            "pack_id": pack.id,
            "language": language,
            "title": (previous or {}).get("title")
            or (out_slides[0]["message"] if out_slides else Path(pptx).stem),
            "base": base_id,
            "slides": out_slides,
        }
        deck = _settle(deck, entries, prev, pack, report)

    for e in entries:
        kind = "drawn" if e.get("drawn") else "imported"
        report[kind].append(e["id"])
        old = prev.get(e["id"])
        if e["spec"] is not None and old is not None and _strip(old) != _strip(e["spec"]):
            report["conflicts"].append(
                {
                    "slide": e["id"],
                    "reason": "changed in Calque after this file was exported: the file wins",
                }
            )
    return deck, extract(base), report


def _strip(s: dict[str, Any]) -> dict[str, Any]:
    return {k: v for k, v in s.items() if v is not None}


def _clone(e: dict[str, Any], prev: dict[str, Any]) -> dict[str, Any]:
    title = " ".join(_title(e["slide"]).split())
    old = e["spec"] or prev.get(e["id"]) or {}
    return {
        "id": e["id"],
        "message": title or old.get("message") or f"Slide {e['n']}",
        "message_type": "imported",
        "form": "imported",
        "source": {"kind": "clone", "from": "base", "slide": e["n"], "values": {}},
    }


def _settle(deck, entries, prev, pack, report) -> dict[str, Any]:
    """Validate; a drawn slide the doctrine refuses where it now sits becomes a clone."""
    by_id = {e["id"]: e for e in entries}
    while True:
        try:
            validate(deck, pack)
            return deck
        except DeckSpecError as err:
            bad = {
                i.slide: i.message
                for i in err.issues
                if i.severity == "ERROR" and i.slide in by_id and by_id[i.slide].get("drawn")
            }
            if not bad:
                raise
            for sid, why in bad.items():
                e = by_id[sid]
                e["drawn"] = None
                report["demoted"].append({"slide": sid, "reason": why})
                deck["slides"] = [_clone(e, prev) if s["id"] == sid else s for s in deck["slides"]]


# --- reference renders -----------------------------------------------------------------------


def _references(specs: list[dict[str, Any]], pack: Pack, language: str, tmp: Path):
    """spec id -> its slide drawn alone on the pack (what the engine wrote), or why it failed."""
    from .build import build

    def run(group: list[dict[str, Any]], name: str):
        out = tmp / f"{name}.pptx"
        deck = {"pack_id": pack.id, "language": language, "title": "reference", "slides": group}
        build(deck, pack, out, check=False)
        return list(Presentation(str(out)).slides)

    if not specs:
        return {}
    try:
        return {s["id"]: sl for s, sl in zip(specs, run(specs, "all"), strict=True)}
    except Exception:  # noqa: BLE001  one bad spec must not cost the others their merge
        pass
    refs: dict[str, Any] = {}
    for i, s in enumerate(specs):
        try:
            refs[s["id"]] = run([s], f"one-{i}")[0]
        except Exception as e:  # noqa: BLE001
            refs[s["id"]] = f"the pack cannot redraw it: {e}"
    return refs


# --- comparison ------------------------------------------------------------------------------


def _paras(tf) -> list[str]:
    return [p.text for p in tf.paragraphs]


_C_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}


def _x(el, path: str) -> list:
    return etree.XPath(path, namespaces=_C_NS)(el)


def _chart_data(chart) -> list[dict[str, Any]]:
    """Series of a chart as its cached values (what PowerPoint saves after a data edit)."""

    def pts(el, path):
        found = _x(el, path)
        if not found:
            return []
        n = _x(found[0], "./c:ptCount/@val")
        size = int(n[0]) if n else 0
        vals: list[str | None] = [None] * size
        for pt in _x(found[0], "./c:pt"):
            i = int(pt.get("idx"))
            v = _x(pt, "./c:v/text()")
            if i >= len(vals):
                vals += [None] * (i + 1 - len(vals))
            vals[i] = v[0] if v else None
        return vals

    out = []
    for ser in _x(chart._chartSpace, ".//c:ser"):
        out.append(
            {
                "name": "".join(_x(ser, "./c:tx//c:v/text()")),
                "cats": pts(ser, "./c:cat//*[c:pt or c:ptCount][1]"),
                "x": pts(ser, "./c:xVal//*[c:pt or c:ptCount][1]"),
                "y": pts(ser, "(./c:val|./c:yVal)//*[c:pt or c:ptCount][1]"),
            }
        )
    return out


_A_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
_RUN_ATTRS = ("sz", "b", "i", "u", "strike", "baseline", "cap")


def _style(el) -> tuple:
    """What a hand edit of a shape's look changes: its fill and line colours, and each run's size,
    weight, colour and face. Run splits and proofing attributes PowerPoint adds do not count."""

    def xp(node, path):
        return etree.XPath(path, namespaces=_A_NS)(node)

    sppr = [x for x in el if x.tag.endswith("}spPr")]
    box = ()
    if sppr:
        box = (
            tuple(xp(sppr[0], "./a:solidFill//@val")),
            tuple(xp(sppr[0], "./a:noFill")) != (),
            tuple(xp(sppr[0], "./a:ln/@w")),
            tuple(xp(sppr[0], "./a:ln/a:solidFill//@val")),
        )
    runs = set()
    for rpr in xp(el, ".//a:r/a:rPr"):
        runs.add(
            (
                *(rpr.get(k) for k in _RUN_ATTRS),
                tuple(xp(rpr, "./a:solidFill//@val")),
                tuple(xp(rpr, "./a:latin/@typeface")),
            )
        )
    return box, tuple(sorted(runs, key=repr))


def _snapshot(slide) -> dict[int, dict[str, Any]]:
    out: dict[int, dict[str, Any]] = {}
    for sh in iter_shapes(slide.shapes):
        snap: dict[str, Any] = {"geom": (sh.left, sh.top, sh.width, sh.height)}
        if sh.shape_type != MSO_SHAPE_TYPE.GROUP:
            snap["style"] = _style(sh._element)
        if getattr(sh, "has_chart", False) and sh.has_chart:
            snap["chart"] = _chart_data(sh.chart)
        elif getattr(sh, "has_table", False) and sh.has_table:
            snap["cells"] = [[c.text_frame.text for c in row.cells] for row in sh.table.rows]
        elif sh.has_text_frame:
            snap["paras"] = _paras(sh.text_frame)
        try:
            snap["image"] = hashlib.sha1(sh.image.blob).hexdigest()
        except (AttributeError, ValueError, KeyError):
            pass
        out[sh.shape_id] = snap
    return out


def _notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def _moved(a, b) -> bool:
    return any(
        (x is None) != (y is None) or (x is not None and abs(x - y) > TOLERANCE_EMU)
        for x, y in zip(a, b, strict=True)
    )


def _merge(
    spec: dict[str, Any], ref, imp, slots: dict[str, int]
) -> tuple[dict[str, Any] | None, str]:
    """The spec with the client's text, notes and chart data edits, or (None, why not)."""
    a, b = _snapshot(ref), _snapshot(imp)
    if set(b) - set(a):
        return None, f"shapes added in PowerPoint ({', '.join(map(str, sorted(set(b) - set(a))))})"
    if set(a) - set(b):
        gone = ", ".join(map(str, sorted(set(a) - set(b))))
        return None, f"shapes removed in PowerPoint ({gone})"
    out = copy.deepcopy(spec)
    for sid, ra in a.items():
        rb = b[sid]
        if _moved(ra["geom"], rb["geom"]):
            return None, f"shape {sid} moved or resized in PowerPoint"
        if ra.get("image") != rb.get("image"):
            return None, f"picture {sid} replaced in PowerPoint"
        if ra.get("style") != rb.get("style"):
            return (
                None,
                f"shape {sid} restyled in PowerPoint (a drawn slide takes its style from the pack)",
            )
        if sid == slots.get("page_number"):
            continue
        if "chart" in ra:
            if ra["chart"] != rb.get("chart"):
                why = _merge_chart(out, rb.get("chart") or [])
                if why:
                    return None, why
        elif "cells" in ra:
            if ra["cells"] != rb.get("cells"):
                why = _merge_cells(out, ra["cells"], rb.get("cells") or [])
                if why:
                    return None, why
        elif "paras" in ra and ra["paras"] != rb.get("paras"):
            new = rb.get("paras") or []
            if sid in (slots.get("title"), slots.get("eyebrow")):
                field = "title" if sid == slots.get("title") else "eyebrow"
                text = "\n".join(new).replace("\v", "\n").strip()
                if field == "title" and not text:
                    return None, "title emptied in PowerPoint"
                out[field] = text or None
                continue
            if len(new) != len(ra["paras"]):
                return None, f"shape {sid}: paragraphs added or removed in PowerPoint"
            for old_p, new_p in zip(ra["paras"], new, strict=True):
                if old_p != new_p and not _replace_leaf(out["source"], old_p, new_p):
                    return None, f"shape {sid}: the edit {old_p!r} -> {new_p!r} has no field"
    notes_a, notes_b = _notes(ref), _notes(imp)
    if notes_a != notes_b:
        out["notes"] = notes_b or None
    return {k: v for k, v in out.items() if v is not None}, ""


def _slots(pack: Pack) -> dict[str, int]:
    """Title / eyebrow / page-number shape ids of the pack's content slide (the drawn canvas)."""
    n = pack.slides_for("content")[0]
    tslide = next(s for s in pack.template_map["slides"] if s["number"] == n)
    out = {k: v for k, v in tslide.get("slots", {}).items() if k in ("title", "eyebrow")}
    if tslide.get("page_number") is not None:
        out["page_number"] = tslide["page_number"]
    return out


def _leaves(obj, path=()):
    if isinstance(obj, str):
        yield path, obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            if k not in ("kind", "id", "type", "chart_type", "icon", "orientation", "split"):
                yield from _leaves(v, (*path, k))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from _leaves(v, (*path, i))


def _set_path(obj, path, value) -> None:
    for k in path[:-1]:
        obj = obj[k]
    obj[path[-1]] = value


def _replace_leaf(source: dict[str, Any], old: str, new: str) -> bool:
    """Write `new` into the one string field the engine drew as `old` (as is, or in caps)."""
    leaves = list(_leaves(source.get("params", {}), ("params",)))
    hits = [p for p, v in leaves if v == old] or [
        p for p, v in leaves if v.upper() == old and v != old
    ]
    if len(hits) != 1:
        return False
    _set_path(source, hits[0], new)
    return True


def _merge_cells(out, old: list[list[str]], new: list[list[str]]) -> str:
    src = out["source"]
    if src.get("id") != "comparison_table":
        return "a table edit on this slide has no field"
    if len(old) != len(new) or any(len(x) != len(y) for x, y in zip(old, new, strict=True)):
        return "table rows or columns added or removed in PowerPoint"
    p = src["params"]
    for r, (ro, rn) in enumerate(zip(old, new, strict=True)):
        for c, (co, cn) in enumerate(zip(ro, rn, strict=True)):
            if co == cn:
                continue
            if r == 0:
                p["header"][c] = cn
            else:
                if r - 1 >= len(p["rows"]):
                    return "the table shows rows the spec does not hold"
                row = p["rows"][r - 1]
                row += [""] * (c + 1 - len(row))
                row[c] = cn
    return ""


def _num(v: str | None) -> float | int | None:
    if v is None:
        return None
    try:
        f = float(v)
    except ValueError:
        return None
    return int(f) if f.is_integer() else f


def _merge_chart(out, data: list[dict[str, Any]]) -> str:
    src = out["source"]
    if src["kind"] == "chart":
        ctype, params = src["type"], src["params"]
    elif src.get("id") == "chart_takeaway":
        ctype, params = src["params"]["chart_type"], src["params"]["chart"]
    else:
        return "chart data edited in PowerPoint, on a slide whose chart has no data field"
    if ctype == "scatter":
        return "scatter data edited in PowerPoint (edit the points in Calque)"
    if not data or not data[0]["cats"]:
        return "chart without categories"
    cats = [c or "" for c in data[0]["cats"]]
    series = []
    for i, d in enumerate(data):
        vals = [_num(v) for v in d["y"]]
        if any(v is None for v in vals) or len(vals) != len(cats):
            return "chart data has empty or non-numeric cells"
        old = params.get("series") or []
        name = d["name"] or (old[i]["name"] if i < len(old) else f"Series {i + 1}")
        series.append({"name": name, "values": vals})
    if ctype == "bar_horizontal":  # drawn bottom-up, listed top-down in the spec
        cats.reverse()
        for s in series:
            s["values"].reverse()
    if ctype == "doughnut":
        series += (params.get("series") or [])[len(series) :]
    params["categories"] = cats
    params["series"] = series
    return ""
