"""Drawing primitives for bespoke slides. Every colour, font and stroke comes from a Style (the
active pack's roles): callers pass role names, never values."""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .slides import crop_to_fill
from .style import Style

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


@dataclass
class Run:
    """One paragraph of a text block. Roles: `font` is a font role, `color` a colour role, `size`
    a size role or points."""

    text: str
    size: str | float = "body"
    color: str = "ink"
    font: str = "body"
    bold: bool = False
    align: str | None = None
    bullet: str | None = None  # a glyph, "num", or None
    level: int = 0
    space_before: float | None = None


def _strip_style(sh):
    """Drop the theme <p:style> python-pptx adds to new shapes: its effectRef re-applies the
    theme's shadow in some renderers. Our shapes style everything explicitly."""
    style = sh._element.find(qn("p:style"))
    if style is not None:
        sh._element.remove(style)
    return sh


def _pt(st: Style, size: str | float) -> float:
    return st.size(size) if isinstance(size, str) else float(size)


def _set_bullet(p, st: Style, run: Run, size_pt: float) -> None:
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buClr", "a:buSzPts", "a:buFont", "a:buChar", "a:buAutoNum", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if run.bullet is None:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        return
    indent = Pt(size_pt * 1.4)
    pPr.set("marL", str(int(indent * (run.level + 1))))
    pPr.set("indent", str(-int(indent)))
    clr = pPr.makeelement(qn("a:buClr"), {})
    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": str(st.rgb(run.color))}))
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buSzPts"), {"val": str(int(size_pt * 100))}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": st.font(run.font)}))
    if run.bullet == "num":
        pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
    else:
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": run.bullet}))


def fill_text(tf, st: Style, runs: str | Run | Sequence[str | Run], align: str = "left") -> None:
    """Fill a text frame: one paragraph per Run (a bare string is a body Run)."""
    items = [runs] if isinstance(runs, (str, Run)) else list(runs)
    for i, item in enumerate(items):
        run = Run(item) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[run.align or align]
        size = _pt(st, run.size)
        if run.space_before:
            p.space_before = Pt(run.space_before)
        r = p.add_run()
        r.text = run.text
        f = r.font
        f.size = Pt(size)
        f.name = st.font(run.font)
        f.bold = run.bold
        f.color.rgb = st.rgb(run.color)
        if run.bullet is not None or i > 0 or len(items) > 1:
            _set_bullet(p, st, run, size)


def add_text(
    slide,
    st: Style,
    l: float,
    t: float,
    w: float,
    h: float,
    runs: str | Run | Sequence[str | Run],
    align: str = "left",
    anchor: str = "top",
):
    """One text box per idea group: a label paragraph then its lines, never one box per line."""
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    fill_text(tf, st, runs, align)
    return tb


def add_box(
    slide,
    st: Style,
    l: float,
    t: float,
    w: float,
    h: float,
    fill: str | None = "surface",
    line: str | None = "rule",
    line_w: str | float = "hairline",
    text: str | Run | Sequence[str | Run] | None = None,
    align: str = "center",
    anchor: str = "middle",
    pad: float = 0.08,
    shape: MSO_SHAPE = MSO_SHAPE.RECTANGLE,
):
    """The diagram unit: a rectangle with optional text. Defaults: surface fill, hairline rule."""
    sh = _strip_style(slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h)))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = st.rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = st.rgb(line)
        sh.line.width = Pt(st.stroke(line_w) if isinstance(line_w, str) else line_w)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(pad * 0.6)
    tf.vertical_anchor = ANCHOR[anchor]
    if text is not None:
        fill_text(tf, st, text, align)
    return sh


def add_line(
    slide,
    st: Style,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    color: str = "muted",
    weight: str | float = "link",
    arrow: str | None = None,
    dash: str | None = None,
):
    """Straight connector; `arrow` = 'end' | 'both' adds triangle heads."""
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = st.rgb(color)
    conn.line.width = Pt(st.stroke(weight) if isinstance(weight, str) else weight)
    ln = conn.line._get_or_add_ln()
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    if arrow in ("end", "both"):
        for tag in ["a:tailEnd"] if arrow == "end" else ["a:headEnd", "a:tailEnd"]:
            ln.append(ln.makeelement(qn(tag), {"type": "triangle", "w": "med", "len": "med"}))
    return _strip_style(conn)


def add_arrow(slide, st: Style, x1: float, y1: float, x2: float, y2: float, **kw):
    kw.setdefault("arrow", "end")
    return add_line(slide, st, x1, y1, x2, y2, **kw)


def add_dot(slide, st: Style, cx: float, cy: float, d: float = 0.11, color: str = "accent"):
    sh = _strip_style(
        slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d)
        )
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = st.rgb(color)
    sh.line.fill.background()
    return sh


def icon_names(st: Style) -> list[str]:
    d = st.pack.dir / "icons"
    return sorted(p.stem for p in d.glob("*.png")) if d.is_dir() else []


def add_icon(
    slide, st: Style, name: str, l: float, t: float, size: float = 0.3, color: str = "ink"
):
    """A pack icon (black on alpha PNG) tinted to a colour role: alpha kept, ink swapped."""
    path = st.pack.dir / "icons" / f"{name}.png"
    if not path.is_file():
        raise KeyError(f"unknown icon {name!r}; pack icons: {', '.join(icon_names(st)) or 'none'}")
    from PIL import Image

    src = Image.open(path).convert("RGBA")
    rgb = st.rgb(color)
    tinted = Image.new("RGBA", src.size, (rgb[0], rgb[1], rgb[2], 0))
    tinted.putalpha(src.getchannel("A"))
    buf = io.BytesIO()
    tinted.save(buf, format="PNG")
    buf.seek(0)
    return slide.shapes.add_picture(buf, Inches(l), Inches(t), Inches(size), Inches(size))


def add_image(slide, path: str | Path, l: float, t: float, w: float, h: float):
    """Image cropped to fill the frame, never distorted."""
    data, _ = crop_to_fill(Path(path).read_bytes(), w / h)
    return slide.shapes.add_picture(io.BytesIO(data), Inches(l), Inches(t), Inches(w), Inches(h))


def _cell_borders(cell, st: Style, bottom: str | None = None, weight: float = 0.75) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    lns = []
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.makeelement(qn(tag), {})
        if tag == "a:lnB" and bottom is not None:
            ln.set("w", str(Pt(weight).emu))
            fill = ln.makeelement(qn("a:solidFill"), {})
            fill.append(ln.makeelement(qn("a:srgbClr"), {"val": str(st.rgb(bottom))}))
            ln.append(fill)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        lns.append(ln)
    for ln in reversed(lns):
        tcPr.insert(0, ln)


def add_table(
    slide,
    st: Style,
    l: float,
    t: float,
    w: float,
    rows: list[list[Any]],
    col_weights: list[float] | None = None,
    header: bool = True,
    size: str | float = "caption",
    row_h: float = 0.32,
    align: list[str] | None = None,
    emphasis: tuple[int, int] | None = None,
):
    """Native table: label-font caps header over an accent hairline, rule hairlines between rows,
    no banding, no fills. `emphasis` = (row, col) of the one cell set in the highlight colour."""
    nrows, ncols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(
        nrows, ncols, Inches(l), Inches(t), Inches(w), Inches(row_h * nrows)
    )
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        for s in tblPr.findall(qn("a:tableStyleId")):
            tblPr.remove(s)
    if col_weights:
        total = float(sum(col_weights))
        for i, cw in enumerate(col_weights):
            tbl.columns[i].width = Inches(w * cw / total)
    pts = _pt(st, size)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(row_h)
        is_header = header and ri == 0
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.background()
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = ALIGN[(align or [])[ci] if align and ci < len(align) else "left"]
            r = p.add_run()
            r.text = str(val).upper() if is_header else str(val)
            f = r.font
            f.name = st.font("label" if is_header else "body")
            f.size = Pt(max(pts - 1, 8) if is_header else pts)
            color = "muted" if is_header else "ink"
            if emphasis == (ri, ci):
                color = "highlight"
            f.color.rgb = st.rgb(color)
            if is_header:
                _cell_borders(cell, st, bottom="accent", weight=st.stroke("link"))
            elif ri < nrows - 1:
                _cell_borders(cell, st, bottom="rule", weight=st.stroke("hairline"))
            else:
                _cell_borders(cell, st)
    return gf


def grid_axes(slide, st: Style, xs: Sequence[float] = (), ys: Sequence[float] = ()):
    """Materialise a few grid axes (3 at most): the grid structures, it never decorates."""
    if len(xs) + len(ys) > 3:
        raise ValueError("at most 3 visible grid axes per slide")
    g = st.grid
    out = [add_line(slide, st, x, g.top, x, g.bottom, color="rule", weight="hairline") for x in xs]
    out += [
        add_line(slide, st, g.margin, y, g.width - g.margin, y, color="rule", weight="hairline")
        for y in ys
    ]
    return out


def item(it: Any) -> tuple[str, str | None, str | None]:
    """DeckSpec Item -> (label, sublabel, icon)."""
    if isinstance(it, str):
        return it, None, None
    return it["label"], it.get("sublabel"), it.get("icon")
