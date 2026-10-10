"""Template extractor: any template.pptx -> template-map + draft DTCG tokens.

Shapes are keyed by `shape_id` (the `<p:cNvPr id>`), never by name: names are copied when a
slide is duplicated and are unreliable. Reviewed keys in an existing map are preserved.
"""

from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

EMU_IN = 914400
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
THEME_SLOTS = (
    "dk1",
    "lt1",
    "dk2",
    "lt2",
    *(f"accent{i}" for i in range(1, 7)),
    "hlink",
    "folHlink",
)
REVIEWED_SLIDE_KEYS = ("description", "slots", "fit", "notes")
PML = "application/vnd.openxmlformats-officedocument.presentationml"
# generic words naming a summary or a closing slide, in the languages packs ship in
SUMMARY = re.compile(
    r"summary|agenda|(table of )?contents|outline|overview|sommaire|ordre du jour|plan", re.I
)
CLOSING = re.compile(r"\b(thanks?|thank you|merci|questions)\b", re.I)
SECTION_NO = re.compile(r"0?\d|\d{2}")  # "01" on a divider


def as_presentation(path: str | Path) -> bool:
    """Rewrite a .potx (template content type) in place as a .pptx. True when it did."""
    path = Path(path)
    with zipfile.ZipFile(path) as z:
        types = z.read("[Content_Types].xml").decode("utf-8")
        if f"{PML}.template.main+xml" not in types:
            return False
        fd, tmp = tempfile.mkstemp(dir=path.parent, suffix=".pptx")
        with open(fd, "wb") as f, zipfile.ZipFile(f, "w", zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = types.replace(
                        f"{PML}.template.main+xml", f"{PML}.presentation.main+xml"
                    ).encode("utf-8")
                out.writestr(item, data)
    shutil.move(tmp, path)
    return True


def used_styles(template: str | Path) -> dict[str, list[str]]:
    """Fonts and hex colours the template's slides set directly, most used first: a real
    template often names its faces on each run, not in the theme."""
    fonts: Counter[str] = Counter()
    colors: Counter[str] = Counter()
    for slide in Presentation(str(template)).slides:
        for el in slide._element.iter():
            face = el.get("typeface")
            if face and not face.startswith("+"):
                fonts[face] += 1
            if el.tag == f"{{{NS['a']}}}srgbClr" and el.get("val"):
                colors[el.get("val").upper()] += 1
    return {
        "fonts": [f for f, _ in fonts.most_common()],
        "colors": [c for c, _ in colors.most_common()],
    }


def _in(v: int | Emu | None) -> float:
    return round((v or 0) / EMU_IN, 2)


def theme(prs) -> dict[str, Any]:
    part = prs.slide_master.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
    )
    root = etree.fromstring(part.blob)
    colors = {}
    for slot in THEME_SLOTS:
        el = root.find(f".//a:clrScheme/a:{slot}", NS)
        if el is None or not len(el):
            continue
        c = el[0]
        # sysClr carries a system name in val; the resolved colour is lastClr
        value = c.get("lastClr") if c.tag.endswith("}sysClr") else c.get("val")
        colors[slot] = (value or "").upper()
    fonts = {}
    for key, tag in (("major", "majorFont"), ("minor", "minorFont")):
        latin = root.find(f".//a:fontScheme/a:{tag}/a:latin", NS)
        if latin is not None:
            fonts[key] = latin.get("typeface")
    return {"colors": colors, "fonts": fonts}


def shape_kind(sh) -> str:
    st = sh.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if getattr(sh, "has_table", False) and sh.has_table:
        return "table"
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return "chart"
    if st == MSO_SHAPE_TYPE.LINE or sh.element.tag.endswith("}cxnSp"):
        return "line"
    if sh.has_text_frame and sh.text_frame.text.strip():
        return "text"
    return "shape"


def _first_run_style(sh) -> tuple[str | None, float | None]:
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip():
                size = r.font.size.pt if r.font.size else None
                return r.font.name, size
    return None, None


def capacity(width_in: float, height_in: float, size_pt: float) -> dict[str, int]:
    # ponytail: average glyph ~0.5 em, line ~1.2 em; lint does real glyph metrics in Phase 2.
    chars = max(1, int(width_in * 72 / (size_pt * 0.5)))
    lines = max(1, round(height_in * 72 / (size_pt * 1.2)))
    return {"chars_per_line": chars, "lines": lines}


def is_page_number(sh, w_in: float, h_in: float) -> bool:
    if not sh.has_text_frame:
        return False
    if sh.element.find(".//a:fld[@type='slidenum']", NS) is not None:
        return True
    text = sh.text_frame.text.strip()
    return (
        bool(re.fullmatch(r"\d{1,3}", text))
        and _in(sh.top) > h_in * 0.85
        and _in(sh.left) > w_in * 0.8
    )


def _shape(sh, w_in: float, h_in: float) -> dict[str, Any]:
    out: dict[str, Any] = {
        "id": sh.shape_id,
        "kind": shape_kind(sh),
        "bbox": [_in(sh.left), _in(sh.top), _in(sh.width), _in(sh.height)],
    }
    if out["kind"] == "text":
        text = sh.text_frame.text.strip()
        out["text"] = text if len(text) <= 80 else text[:77] + "..."
        font, size = _first_run_style(sh)
        if font:
            out["font"] = font
        if size:
            out["size"] = size
            out["capacity"] = capacity(out["bbox"][2], out["bbox"][3], size)
    elif out["kind"] == "table":
        out["table"] = [len(sh.table.rows), len(sh.table.columns)]
    elif out["kind"] == "group":
        out["children"] = [_shape(c, w_in, h_in) for c in sh.shapes]
    return out


def extract(template: str | Path, previous: dict[str, Any] | None = None) -> dict[str, Any]:
    prs = Presentation(str(template))
    w_in, h_in = _in(prs.slide_width), _in(prs.slide_height)
    prev = {s["number"]: s for s in (previous or {}).get("slides", [])}
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        entry: dict[str, Any] = {"number": i, "layout": slide.slide_layout.name}
        for key in REVIEWED_SLIDE_KEYS:  # reviewed keys first: what a reviewer reads
            if key in prev.get(i, {}):
                entry[key] = prev[i][key]
        page = [sh.shape_id for sh in slide.shapes if is_page_number(sh, w_in, h_in)]
        entry["page_number"] = page[0] if page else None
        entry["shapes"] = [_shape(sh, w_in, h_in) for sh in slide.shapes]
        slides.append(entry)
    return {
        "canvas": {"width_in": w_in, "height_in": h_in},
        "theme": theme(prs),
        "slides": slides,
    }


def _text_shapes(slides):
    for s in slides:
        stack = list(s["shapes"])
        while stack:
            sh = stack.pop()
            stack.extend(sh.get("children", []))
            if sh["kind"] == "text" and sh.get("size"):
                yield s, sh


def draft_tokens(tmap: dict[str, Any]) -> dict[str, Any]:
    """A first tokens.json from the template theme and observed type sizes. Meant for review."""
    th = tmap["theme"]
    h_in = tmap["canvas"]["height_in"]
    shapes = list(_text_shapes(tmap["slides"]))
    sizes = Counter(round(sh["size"]) for _, sh in shapes)
    body = sizes.most_common(1)[0][0] if sizes else 12
    tops = Counter(round(sh["size"]) for _, sh in shapes if sh["bbox"][1] < h_in * 0.15)
    title = tops.most_common(1)[0][0] if tops else round(body * 2)
    biggest = max(sizes) if sizes else title * 3
    cover = max((sh["size"] for s, sh in shapes if s["number"] == 1), default=title * 2)

    def color(slot: str) -> dict[str, str]:
        return {"$value": f"{{theme.{slot}}}"}

    return {
        "theme": {
            "$type": "color",
            **{slot: {"$value": f"#{hexv.lower()}"} for slot, hexv in th["colors"].items()},
            "font": {
                "$type": "fontFamily",
                **{k: {"$value": v} for k, v in th["fonts"].items()},
            },
        },
        "role": {
            "color": {
                "$type": "color",
                "ink": color("dk1"),
                "background": color("lt1"),
                "accent": color("accent1"),
                "highlight": color("accent2"),
                "muted": color("dk2"),
                "rule": color("lt2"),
                "surface": color("lt1"),
                "series": {"1": color("accent3"), "2": color("accent4"), "3": color("accent5")},
                "status": {
                    "positive": color("accent6"),
                    "warning": color("accent4"),
                    "info": color("accent5"),
                },
            },
            "font": {
                "$type": "fontFamily",
                "display": {"$value": "{theme.font.major}"},
                "body": {"$value": "{theme.font.minor}"},
                "label": {"$value": "{theme.font.minor}"},
            },
            "fontWeight": {
                "$type": "fontWeight",
                "display": {"$value": 400},
                "body": {"$value": 400},
                "label": {"$value": 400},
            },
            "size": {
                "$type": "dimension",
                **{
                    k: {"$value": f"{v}pt"}
                    for k, v in {
                        "cover_title": round(cover),
                        "divider_title": round(cover * 0.85),
                        "title": title,
                        "body": body,
                        "label": max(8, round(body * 0.75)),
                        "caption": max(8, round(body * 0.75)),
                        "numeral": round(biggest * 0.6),
                        "numeral_xl": round(biggest),
                    }.items()
                },
            },
            "stroke": {
                "$type": "dimension",
                "hairline": {"$value": "0.75pt"},
                "link": {"$value": "1pt"},
                "emphasis": {"$value": "1.5pt"},
            },
        },
    }


def _words(sh: dict[str, Any]) -> int:
    return len(sh["text"].split())


def _is_divider(s: dict[str, Any]) -> bool:
    """A section break: named so by its layout, or a few short texts led by a section number."""
    if re.search(r"section|divider|chapter", s["layout"], re.I):
        return True
    texts = [sh for sh in _texts(s) if sh["id"] != s.get("page_number")]
    if not texts or len(texts) > 8:
        return False
    biggest = max(texts, key=lambda sh: sh.get("size") or 0)
    numbered = any(SECTION_NO.fullmatch(sh["text"]) for sh in texts)
    return numbered and _words(biggest) <= 6 and bool(re.search(r"[^\W\d_]", biggest["text"]))


def _role_numbers(slides: list[dict[str, Any]], h_in: float) -> dict[str, list[int]]:
    """Guess roles from slide order, layout names, text structure and a few generic words
    ("Agenda", "Thank you"); the reviewer corrects them."""
    cover = slides[0]["number"]
    rest = slides[1:]
    # a closing says thanks in a few words; else the last slide
    closing = [
        s["number"]
        for s in rest
        if 0 < len(_texts(s)) <= 4 and any(CLOSING.search(sh["text"]) for sh in _texts(s))
    ] or [slides[-1]["number"]]
    rest = [s for s in rest if s["number"] not in closing]
    divider = [s["number"] for s in rest if _is_divider(s)]
    titles = {
        max(_texts(s), key=lambda sh: sh.get("size") or 0)["text"].lower()
        for s in rest
        if s["number"] in divider and _texts(s)
    }
    # a summary is titled so, or lists the titles of at least two dividers
    summary = [
        s["number"]
        for s in rest
        if s["number"] not in divider
        and (
            any(SUMMARY.fullmatch(sh["text"].strip(" :.")) for sh in _texts(s))
            or len(titles & {sh["text"].lower() for sh in _texts(s)}) >= 2
        )
    ]
    rest = [s for s in rest if s["number"] not in divider + summary]
    content = [
        s["number"] for s in rest if _texts(s) and re.search(r"title only|blank", s["layout"], re.I)
    ]
    if not content:
        # the plainest slide whose biggest text is a title at the top: a canvas to compose on
        titled = [
            s
            for s in rest
            if _texts(s)
            and max(_texts(s), key=lambda sh: sh.get("size") or 0)["bbox"][1] < h_in * 0.25
        ]
        content = [min(titled or rest or slides, key=lambda s: len(s["shapes"]))["number"]]
    roles = {"cover": [cover], "content": content, "closing": closing}
    if summary:
        roles["summary"] = summary
    if divider:
        roles["divider"] = divider
    return roles


def _texts(slide: dict[str, Any]) -> list[dict[str, Any]]:
    return [sh for sh in slide["shapes"] if sh["kind"] == "text" and sh.get("text")]


def draft_manifest(
    tmap: dict[str, Any], pack_id: str, name: str, used: dict[str, list[str]] | None = None
) -> dict[str, Any]:
    """A first pack.yaml for an imported template, and title/subtitle slots on its role slides
    (written into `tmap`). Meant for review: roles, grid and placeholders are guesses.
    `used`: the template's fonts and colours (`used_styles`), allowed by lint."""
    w_in, h_in = tmap["canvas"]["width_in"], tmap["canvas"]["height_in"]
    by_num = {s["number"]: s for s in tmap["slides"]}
    roles = _role_numbers(tmap["slides"], h_in)
    role_slides = sorted({n for ns in roles.values() for n in ns})
    for n in role_slides:
        s = by_num[n]
        if "slots" in s:
            continue
        texts = sorted(_texts(s), key=lambda sh: (-(sh.get("size") or 0), sh["bbox"][1]))
        page = s.get("page_number")
        # not the page number, a section number or a footer
        texts = [
            sh
            for sh in texts
            if sh["id"] != page
            and not SECTION_NO.fullmatch(sh["text"])
            and sh["bbox"][1] < h_in * 0.85
        ]
        if texts:
            s["slots"] = {"title": texts[0]["id"]}
            # a summary's entries are reviewed into slots (ch1_title...), not guessed
            if len(texts) > 1 and n not in roles.get("summary", []):
                s["slots"]["subtitle"] = texts[1]["id"]

    content = by_num[roles["content"][0]]
    tops = sorted(_texts(content) or content["shapes"], key=lambda sh: sh["bbox"][1])
    left, top, width, height = tops[0]["bbox"] if tops else [0.5, 0.3, w_in - 1.0, 1.0]
    margin, gutter = round(left, 2), 0.3

    def columns(n: int) -> list[float]:
        col = (w_in - 2 * margin - (n - 1) * gutter) / n
        return [round(margin + i * (col + gutter), 2) for i in range(n)]

    grid: dict[str, Any] = {
        "margin_in": margin,
        "columns": {"3": columns(3), "4": columns(4)},
        "gutter_in": gutter,
        "title": {"left_in": left, "top_in": top, "width_in": width, "height_in": height},
        "body_top_in": round(top + height + 0.25, 2),
        "footer_top_in": round(h_in - 0.6, 2),
    }
    page = next(
        (
            sh
            for s in tmap["slides"]
            for sh in s["shapes"]
            if s.get("page_number") and sh["id"] == s["page_number"]
        ),
        None,
    )
    if page:
        grid["page_number"] = {"left_in": page["bbox"][0], "top_in": page["bbox"][1]}

    # sample copy in the slots a deck fills ("Presentation title") must not survive in a deck.
    # Single words, the closing slide's text ("Contact", "Thank you") and text a clone keeps
    # (a footer, a summary's entries) are real copy, not placeholders.
    slotted = {
        n: set(by_num[n].get("slots", {}).values())
        for n in role_slides
        if n not in roles["closing"]
    }
    kept = [
        sh["text"].lower()
        for n in role_slides
        for sh in _texts(by_num[n])
        if sh["id"] not in slotted.get(n, ())
    ]
    placeholders = sorted(
        {
            sh["text"]
            for n, ids in slotted.items()
            for sh in _texts(by_num[n])
            if sh["id"] in ids
            and _words(sh) >= 2
            and not sh["text"].endswith("...")
            and not any(sh["text"].lower() in k for k in kept)
        }
    )
    fonts = tmap["theme"]["fonts"]
    theme_colors = set(tmap["theme"]["colors"].values())
    lint: dict[str, Any] = {"placeholders": placeholders}
    if used:  # the template's own faces and colours are the charter
        extra_fonts = [f for f in used["fonts"] if f not in fonts.values()]
        extra_colors = [c for c in used["colors"] if c not in theme_colors]
        if extra_fonts:
            lint["extra_fonts"] = extra_fonts
        if extra_colors:
            lint["extra_colors"] = extra_colors
    return {
        "id": pack_id,
        "name": name,
        "version": "0.1.0",
        "default_language": "en",
        "missing_value": {"en": "[TO COMPLETE]", "fr": "[À COMPLÉTER]"},
        "roles": roles,
        "never_clone": [],
        "grid": grid,
        "fonts": {
            "fallback": {
                "display": fonts.get("major", "sans-serif"),
                "body": fonts.get("minor", "sans-serif"),
                "label": fonts.get("minor", "sans-serif"),
            }
        },
        "lint": lint,
    }
